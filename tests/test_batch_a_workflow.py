import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from forms.engineering.form_2_7 import FORM_2_7
from forms.engineering.form_2_8 import FORM_2_8
from forms.engineering.form_2_9 import FORM_2_9
from forms.engineering.form_2_12 import FORM_2_12
from forms.engineering.form_2_13 import FORM_2_13

from forms.engineering.persistence import (
    complete_engineering_record,
    create_engineering_record,
    get_engineering_record,
    update_engineering_record,
)

from services.engineering_original_form_export import (
    export_engineering_original_form,
)
from services.engineering_summary_export import (
    export_engineering_summary,
)

from tests.workflow_test_support import (
    EngineeringWorkflowTestCaseBase,
)


class BatchAWorkflowMixin:
    """
    Batch A 五张新增工程调查表的统一生产级 workflow 合同。

    每张表都验证：
    - 新增草稿；
    - point / range 工程位置；
    - 同位置重复保护；
    - 完成调查；
    - completed 后继续修改；
    - 当前批次 / DataQuery 共用查询结果；
    - 详细汇总实际导出；
    - 正式原表实际导出及全部声明式绑定。
    """

    DEFINITION = None
    ASSET_NAME = None

    POINT_STAKE = "K70+500"
    POINT_STAKE_VALUE = 70500.0

    RANGE_START = "K70+100"
    RANGE_START_VALUE = 70100.0
    RANGE_END = "K70+900"
    RANGE_END_VALUE = 70900.0

    def setUp(self):
        super().setUp()

        if self.DEFINITION is None:
            raise ValueError(
                "Batch A workflow 必须配置 DEFINITION。"
            )

        self.assertEqual(
            self.DEFINITION.form_code,
            self.FORM_CODE,
        )

    # =========================================================
    # payload helpers
    # =========================================================

    def _position(self):
        if self.DEFINITION.position.kind == "point":
            return {
                "kind": "point",
                "single_stake_text": self.POINT_STAKE,
                "single_stake_value": self.POINT_STAKE_VALUE,
            }

        return {
            "kind": "range",
            "start_stake_text": self.RANGE_START,
            "start_stake_value": self.RANGE_START_VALUE,
            "end_stake_text": self.RANGE_END,
            "end_stake_value": self.RANGE_END_VALUE,
        }

    def _record_data(self, *, asset_name=None):
        asset_name = asset_name or self.ASSET_NAME

        result = {}

        for index, field in enumerate(
            self.DEFINITION.fields,
            start=1,
        ):
            if field.key == self.DEFINITION.asset_name_field:
                value = asset_name

            elif field.input_type == "stake":
                if field.key == "start_stake":
                    value = self.RANGE_START
                elif field.key == "end_stake":
                    value = self.RANGE_END
                else:
                    value = self.POINT_STAKE

            elif field.input_type == "month":
                value = (
                    "2024-01"
                    if not field.required
                    else "2020-06"
                )

            elif field.input_type == "choice":
                self.assertTrue(field.choices)
                value = field.choices[0]

            elif field.input_type == "integer":
                value = index + 1

            elif field.input_type == "signed_decimal":
                value = -(
                    1000.0
                    + index
                    + 0.25
                )

            elif field.input_type == "decimal":
                value = (
                    10.0
                    + index
                    + 0.25
                )

            elif field.input_type == "text":
                value = (
                    f"测试{field.label}"
                )

            else:
                raise AssertionError(
                    "测试尚未覆盖字段类型："
                    f"{field.input_type}"
                )

            result[field.key] = value

        return result

    def _inspection_results(self, grade="B"):
        return [
            {
                "item_code": item["item_code"],
                "category": item["category"],
                "item_name": item["item_name"],
                "grade": grade,
                "description": None,
                "remark": None,
            }
            for item
            in self.DEFINITION.evaluation_items
        ]

    def _payload(
        self,
        *,
        asset_name=None,
        survey_comment="Batch A workflow 测试。",
    ):
        asset_name = asset_name or self.ASSET_NAME

        return {
            "asset_name": asset_name,
            "record_data": self._record_data(
                asset_name=asset_name
            ),
            "position": self._position(),
            "inspection_results": (
                self._inspection_results()
            ),
            "survey_date": "2026-09-17",
            "overall_grade": "B",
            "survey_comment": survey_comment,
        }

    def _business_code(self, sequence=1):
        return (
            "1-01-01-"
            f"{self.DEFINITION.business_type_code}"
            f"-{sequence:03d}"
        )

    def _create_draft(
        self,
        *,
        sequence=1,
        asset_name=None,
    ):
        return create_engineering_record(
            self.DEFINITION,
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=int(
                self.form_version["id"]
            ),
            organization_unit_id=self.office_id,
            canal_unit_id=self.canal_id,
            business_code=self._business_code(
                sequence
            ),
            payload=self._payload(
                asset_name=asset_name
            ),
        )

    @staticmethod
    def _resolve_original_binding(
        *,
        record,
        binding,
    ):
        record_data = (
            record["record_data"]
            or {}
        )

        if binding.source == "record":
            values = [
                record[key]
                for key in binding.keys
            ]

        elif binding.source == "record_data":
            values = [
                record_data.get(key)
                for key in binding.keys
            ]

        else:
            raise AssertionError(
                "正式原表测试出现未支持数据源："
                f"{binding.source}"
            )

        if binding.formatter is not None:
            return binding.formatter(
                *values
            )

        return values[0]

    # =========================================================
    # 1. 草稿 / 列表 / 重复保护
    # =========================================================

    def test_draft_query_and_duplicate_protection(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result["survey_record_id"]
        )

        record = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(record)
        assert record is not None

        self.assertEqual(
            record["record_status"],
            "draft",
        )
        self.assertEqual(
            record["asset_type"],
            self.DEFINITION.asset_type,
        )
        self.assertEqual(
            record["asset_name"],
            self.ASSET_NAME,
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=self.project_id,
                survey_batch_id=self.batch_id,
                form_code=(
                    self.DEFINITION
                    .form_code
                ),
            )
        )

        self.assertEqual(
            len(records),
            1,
        )
        self.assertEqual(
            records[0]["record_status"],
            "draft",
        )

        position_text = (
            records[0][
                "engineering_position"
            ]
            or ""
        )

        if (
            self.DEFINITION.position.kind
            == "point"
        ):
            self.assertIn(
                self.POINT_STAKE,
                position_text,
            )
        else:
            self.assertIn(
                self.RANGE_START,
                position_text,
            )
            self.assertIn(
                self.RANGE_END,
                position_text,
            )

        with self.assertRaises(
            ValueError
        ):
            self._create_draft(
                sequence=2,
                asset_name=(
                    self.ASSET_NAME
                    + "重复"
                ),
            )

    # =========================================================
    # 2. 完成 -> completed 修改
    # =========================================================

    def test_complete_then_edit_completed_record(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result["survey_record_id"]
        )

        completion = (
            complete_engineering_record(
                self.DEFINITION,
                survey_record_id=(
                    survey_record_id
                ),
                payload=self._payload(),
            )
        )

        self.assertEqual(
            completion[
                "inspection_count"
            ],
            len(
                self.DEFINITION
                .evaluation_items
            ),
        )

        completed = (
            get_engineering_record(
                self.DEFINITION,
                survey_record_id=(
                    survey_record_id
                ),
            )
        )

        self.assertIsNotNone(completed)
        assert completed is not None

        self.assertEqual(
            completed["record_status"],
            "completed",
        )
        self.assertEqual(
            completed["overall_grade"],
            "B",
        )

        changed_name = (
            self.ASSET_NAME
            + "（已完成修改）"
        )

        update_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(
                asset_name=changed_name,
                survey_comment=(
                    "completed 后修改成功。"
                ),
            ),
        )

        updated = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(updated)
        assert updated is not None

        self.assertEqual(
            updated["record_status"],
            "completed",
        )
        self.assertEqual(
            updated["asset_name"],
            changed_name,
        )
        self.assertEqual(
            updated["survey_comment"],
            "completed 后修改成功。",
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=self.project_id,
                survey_batch_id=self.batch_id,
                form_code=(
                    self.DEFINITION
                    .form_code
                ),
            )
        )

        self.assertEqual(
            len(records),
            1,
        )
        self.assertEqual(
            records[0]["overall_grade"],
            "B",
        )
        self.assertEqual(
            records[0]["record_status"],
            "completed",
        )

    # =========================================================
    # 3. 详细汇总实际导出
    # =========================================================

    def test_summary_export_real_workbook(
        self,
    ):
        result = self._create_draft()
        survey_record_id = int(
            result["survey_record_id"]
        )

        complete_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(),
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=self.project_id,
                survey_batch_id=self.batch_id,
                form_code=(
                    self.DEFINITION
                    .form_code
                ),
            )
        )

        output = (
            Path(
                self.temp_directory.name
            )
            / (
                self.DEFINITION.form_code
                + "_summary.xlsx"
            )
        )

        export_result = (
            export_engineering_summary(
                self.DEFINITION,
                records=records,
                file_path=output,
            )
        )

        self.assertEqual(
            export_result[
                "exported_count"
            ],
            1,
        )
        self.assertTrue(
            output.exists()
        )

        workbook = load_workbook(
            output,
            data_only=False,
        )

        try:
            summary = (
                self.DEFINITION
                .summary_export_definition
            )
            self.assertIsNotNone(
                summary
            )
            assert summary is not None

            worksheet = workbook[
                summary.sheet_name
            ]

            headers = {
                cell.value: cell.column
                for cell in worksheet[1]
            }

            self.assertIn(
                "业务编号",
                headers,
            )
            self.assertIn(
                "名称",
                headers,
            )
            self.assertIn(
                "工程状况类别",
                headers,
            )
            self.assertIn(
                "状态",
                headers,
            )

            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "业务编号"
                    ],
                ).value,
                self._business_code(),
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "名称"
                    ],
                ).value,
                self.ASSET_NAME,
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "工程状况类别"
                    ],
                ).value,
                "B",
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "状态"
                    ],
                ).value,
                "录入完成",
            )

            for item in (
                self.DEFINITION
                .evaluation_items
            ):
                header = (
                    f"{item['category']}"
                    f"-{item['item_name']}"
                )
                self.assertIn(
                    header,
                    headers,
                )
                self.assertEqual(
                    worksheet.cell(
                        row=2,
                        column=headers[
                            header
                        ],
                    ).value,
                    "B",
                )

        finally:
            workbook.close()

    # =========================================================
    # 4. 正式原表实际导出
    # =========================================================

    def test_original_form_real_export(
        self,
    ):
        result = self._create_draft()
        survey_record_id = int(
            result["survey_record_id"]
        )

        complete_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(),
        )

        original = (
            self.DEFINITION
            .original_form_export_definition
        )

        self.assertIsNotNone(original)
        assert original is not None

        output = (
            Path(
                self.temp_directory.name
            )
            / (
                self.DEFINITION.form_code
                + "_original.xlsx"
            )
        )

        export_engineering_original_form(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            file_path=output,
        )

        self.assertTrue(
            output.exists()
        )

        record = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(record)
        assert record is not None

        workbook = load_workbook(
            output,
            data_only=False,
        )

        try:
            worksheet = workbook[
                original.sheet_name
            ]

            # 顶部归属公共区。
            self.assertEqual(
                worksheet["B3"].value,
                "处",
            )
            self.assertEqual(
                worksheet["D3"].value,
                "所",
            )
            self.assertEqual(
                worksheet["F3"].value,
                "干渠",
            )
            self.assertEqual(
                worksheet["H3"].value,
                "支渠",
            )
            self.assertEqual(
                worksheet["I3"].value,
                "编号：",
            )
            self.assertEqual(
                worksheet["J3"].value,
                self._business_code(),
            )

            # Definition 中声明的每一个正式原表字段
            # 都必须真实写入模板指定单元格。
            for cell_binding in (
                original.field_bindings
            ):
                expected = (
                    self
                    ._resolve_original_binding(
                        record=record,
                        binding=(
                            cell_binding
                            .binding
                        ),
                    )
                )

                actual = worksheet[
                    cell_binding.cell
                ].value

                self.assertEqual(
                    actual,
                    expected,
                    msg=(
                        self.DEFINITION
                        .form_code
                        + " "
                        + cell_binding.cell
                    ),
                )

            # 全部分项评价。
            start_row = (
                original
                .evaluation_binding
                .start_row
            )
            grade_column = (
                original
                .evaluation_binding
                .column
            )

            for offset in range(
                len(
                    self.DEFINITION
                    .evaluation_items
                )
            ):
                self.assertEqual(
                    worksheet[
                        f"{grade_column}"
                        f"{start_row + offset}"
                    ].value,
                    "B",
                )

            conclusion = (
                original
                .conclusion_binding
            )

            self.assertEqual(
                worksheet[
                    conclusion
                    .survey_comment_cell
                ].value,
                "Batch A workflow 测试。",
            )
            self.assertEqual(
                worksheet[
                    conclusion
                    .overall_grade_cell
                ].value,
                "B",
            )
            self.assertEqual(
                worksheet[
                    conclusion
                    .survey_date_cell
                ].value,
                "2026-09-17",
            )

            self.assertEqual(
                worksheet.print_area
                .replace("$", ""),
                (
                    f"'{original.sheet_name}'!"
                    f"{original.print_settings.print_area}"
                ),
            )

        finally:
            workbook.close()


class Form27WorkflowTestCase(
    BatchAWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_7"
    DEFINITION = FORM_2_7
    TEST_DB_FILENAME = (
        "test_form_2_7_batch_a.db"
    )
    ASSET_NAME = "测试跌水与陡坡"


class Form28WorkflowTestCase(
    BatchAWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_8"
    DEFINITION = FORM_2_8
    TEST_DB_FILENAME = (
        "test_form_2_8_batch_a.db"
    )
    ASSET_NAME = "测试渠下涵"


class Form29WorkflowTestCase(
    BatchAWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_9"
    DEFINITION = FORM_2_9
    TEST_DB_FILENAME = (
        "test_form_2_9_batch_a.db"
    )
    ASSET_NAME = "测试闸门及启闭设施"


class Form212WorkflowTestCase(
    BatchAWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_12"
    DEFINITION = FORM_2_12
    TEST_DB_FILENAME = (
        "test_form_2_12_batch_a.db"
    )
    ASSET_NAME = "测试桥梁"


class Form213WorkflowTestCase(
    BatchAWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_13"
    DEFINITION = FORM_2_13
    TEST_DB_FILENAME = (
        "test_form_2_13_batch_a.db"
    )
    ASSET_NAME = "测试砌石工程"


if __name__ == "__main__":
    unittest.main()
