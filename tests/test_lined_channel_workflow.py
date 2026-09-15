import json
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

# =========================
# 让测试可以导入 src
# =========================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from tests.workflow_test_support import (
    EngineeringWorkflowTestCaseBase,
)

from services.lined_channel_evaluation import (
    LINED_CHANNEL_EVALUATION_ITEMS,
)
from services.lined_channel_export import (
    export_lined_channel_original_form,
    export_lined_channel_summary,
)


class LinedChannelWorkflowTestCase(
    EngineeringWorkflowTestCaseBase,
):
    """
    附表2.1防渗衬砌渠道渠段
    V0.3.0-A1 最小持久化测试。

    每个测试使用独立临时 SQLite 数据库，
    不接触正式数据库。
    """

    FORM_CODE = "form_2_1"

    TEST_DB_FILENAME = "test_yinda_survey.db"

    PROJECT_NAME = "附表2.1自动测试项目"
    PROJECT_SHORT_NAME = "2.1测试"

    BATCH_NAME = "附表2.1自动测试批次"
    BATCH_CODE = "FORM_2_1_TEST"

    BATCH_START_DATE = "2026-09-01"
    BATCH_END_DATE = "2026-12-31"

    DEPARTMENT_NAME = "测试基层处"
    DEPARTMENT_CODE = "1"

    OFFICE_NAME = "测试水管所"
    OFFICE_CODE = "01"

    CANAL_NAME = "测试总干渠"
    CANAL_LEVEL = "01"
    CANAL_DESCRIPTION = None

    # =========================
    # 测试1：
    # 附表2.1元数据正确存在
    # =========================

    def test_form_2_1_definition_exists(self):
        form_version = database.get_current_form_version("form_2_1")

        self.assertIsNotNone(form_version)

        assert form_version is not None

        self.assertEqual(
            form_version["form_number"],
            "2.1",
        )

        self.assertEqual(
            form_version["asset_type"],
            "lined_channel_section",
        )

    # =========================
    # 测试2：
    # 渠段型 EngineeringAsset
    # 正确保存起止桩号
    # =========================

    def test_create_lined_channel_section_survey(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        self.assertIsNotNone(form_version)

        assert form_version is not None

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=form_version["id"],
            asset_name="测试总干渠",
            organization_unit_id=self.office_id,
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": "测试总干渠",
                "section_length": 1250.0,
            },
            start_stake_text="K12+000",
            start_stake_value=12000.0,
            end_stake_text="K13+250",
            end_stake_value=13250.0,
        )

        engineering_asset_id = int(result["engineering_asset_id"])

        survey_record_id = int(result["survey_record_id"])

        # -------------------------
        # 1. 检查 EngineeringAsset
        # -------------------------

        asset = database.get_engineering_asset_detail(engineering_asset_id)

        self.assertIsNotNone(asset)

        assert asset is not None

        self.assertEqual(
            asset["asset_type"],
            "lined_channel_section",
        )

        # 渠段型工程不能写入单桩号
        self.assertIsNone(asset["single_stake_text"])

        self.assertEqual(
            asset["start_stake_text"],
            "K12+000",
        )

        self.assertEqual(
            float(asset["start_stake_value"]),
            12000.0,
        )

        self.assertEqual(
            asset["end_stake_text"],
            "K13+250",
        )

        self.assertEqual(
            float(asset["end_stake_value"]),
            13250.0,
        )

        # -------------------------
        # 2. 检查 SurveyRecord
        # -------------------------

        with database.get_connection() as connection:
            record = connection.execute(
                """
                SELECT
                    project_id,
                    survey_batch_id,
                    form_version_id,
                    record_type,
                    engineering_asset_id,
                    business_code,
                    record_status,
                    record_data_json
                FROM survey_records
                WHERE id = ?
                """,
                (survey_record_id,),
            ).fetchone()

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            int(record["project_id"]),
            self.project_id,
        )

        self.assertEqual(
            int(record["survey_batch_id"]),
            self.batch_id,
        )

        self.assertEqual(
            int(record["form_version_id"]),
            int(form_version["id"]),
        )

        self.assertEqual(
            int(record["engineering_asset_id"]),
            engineering_asset_id,
        )

        self.assertEqual(
            record["record_type"],
            "engineering",
        )

        self.assertEqual(
            record["record_status"],
            "draft",
        )

        self.assertEqual(
            record["business_code"],
            "1-01-01-01-001",
        )

        record_data = json.loads(record["record_data_json"])

        self.assertEqual(
            record_data["channel_name"],
            "测试总干渠",
        )

        self.assertEqual(
            record_data["section_length"],
            1250.0,
        )

        # -------------------------
        # 3. 检查现有工程台账查询
        # -------------------------

        assets = database.get_engineering_assets(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
        )

        self.assertEqual(
            len(assets),
            1,
        )

        ledger_asset = assets[0]

        self.assertEqual(
            ledger_asset["start_stake_text"],
            "K12+000",
        )

        self.assertEqual(
            ledger_asset["end_stake_text"],
            "K13+250",
        )

        self.assertIsNone(ledger_asset["single_stake_text"])

    def test_lined_channel_record_can_be_loaded(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=form_version["id"],
            asset_name="测试总干渠",
            organization_unit_id=self.office_id,
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": "测试总干渠",
                "section_length": 1250.0,
                "build_date": "2008-06",
            },
            start_stake_text="K12+000",
            start_stake_value=12000.0,
            end_stake_text="K13+250",
            end_stake_value=13250.0,
        )

        record = database.get_lined_channel_section_record(result["survey_record_id"])

        self.assertEqual(
            record["asset_name"],
            "测试总干渠",
        )

        self.assertEqual(
            record["asset_type"],
            "lined_channel_section",
        )

        self.assertEqual(
            record["start_stake_text"],
            "K12+000",
        )

        self.assertEqual(
            record["end_stake_text"],
            "K13+250",
        )

        self.assertEqual(
            record["record_status"],
            "draft",
        )

        self.assertEqual(
            record["record_data"]["build_date"],
            "2008-06",
        )

    def test_lined_channel_draft_can_be_updated(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=form_version["id"],
            asset_name="测试总干渠",
            organization_unit_id=self.office_id,
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": "测试总干渠",
                "section_length": 1250.0,
            },
            start_stake_text="K12+000",
            start_stake_value=12000.0,
            end_stake_text="K13+250",
            end_stake_value=13250.0,
        )

        survey_record_id = result["survey_record_id"]

        engineering_asset_id = result["engineering_asset_id"]

        database.update_lined_channel_section_draft(
            survey_record_id=survey_record_id,
            asset_name="测试总干渠修改后",
            organization_unit_id=self.office_id,
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": ("测试总干渠修改后"),
                "section_length": 1500.0,
                "build_date": "2010-08",
            },
            start_stake_text="K12+100",
            start_stake_value=12100.0,
            end_stake_text="K13+600",
            end_stake_value=13600.0,
        )

        record = database.get_lined_channel_section_record(survey_record_id)

        # 仍然是原来的 EngineeringAsset
        self.assertEqual(
            record["engineering_asset_id"],
            engineering_asset_id,
        )

        # 仍然是原来的 SurveyRecord
        self.assertEqual(
            record["survey_record_id"],
            survey_record_id,
        )

        self.assertEqual(
            record["asset_name"],
            "测试总干渠修改后",
        )

        self.assertEqual(
            record["start_stake_text"],
            "K12+100",
        )

        self.assertEqual(
            record["end_stake_text"],
            "K13+600",
        )

        self.assertIsNone(record["single_stake_text"])

        self.assertEqual(
            record["record_data"]["section_length"],
            1500.0,
        )

        self.assertEqual(
            record["record_data"]["build_date"],
            "2010-08",
        )

        records = database.get_lined_channel_section_records(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
        )

        self.assertEqual(
            len(records),
            1,
        )

        self.assertEqual(
            records[0]["survey_record_id"],
            survey_record_id,
        )

        self.assertEqual(
            records[0]["start_stake_text"],
            "K12+100",
        )

        self.assertEqual(
            records[0]["end_stake_text"],
            "K13+600",
        )

    def test_all_basic_fields_round_trip(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        record_data = {
            "channel_name": "测试一干渠",
            "start_stake": "K10+000",
            "start_stake_value": 10000.0,
            "end_stake": "K11+500",
            "end_stake_value": 11500.0,
            "section_length": 1500.0,
            "build_date": "2008-06",
            "renovation_date": "2021-09",
            "longitudinal_slope": "1/2000",
            "design_flow": 12.5,
            "channel_grade": "3级",
            "cross_section_form": "梯形",
            "embankment_top_width": 3.5,
            "inner_slope": "1:1.5",
            "outer_slope": "1:1.5",
            "increased_flow": 15.0,
            "bed_soil": "砂壤土",
            "lining_structure": ("现浇混凝土衬砌"),
            "freeboard": 0.6,
            "bottom_width": 4.2,
            "water_conveyance_loss": 0.08,
            "lining_material": "混凝土",
            "lining_thickness": 12.0,
            "concrete_strength": "C25",
            "channel_depth": 2.8,
            "channel_bottom_elevation": (1865.35),
        }

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试一干渠",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data=record_data,
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+500",
            end_stake_value=11500.0,
        )

        loaded = database.get_lined_channel_section_record(result["survey_record_id"])

        loaded_data = loaded["record_data"]

        self.assertEqual(
            loaded["start_stake_text"],
            "K10+000",
        )

        self.assertEqual(
            loaded["end_stake_text"],
            "K11+500",
        )

        for key, expected_value in record_data.items():
            self.assertEqual(
                loaded_data[key],
                expected_value,
                msg=f"字段回填失败：{key}",
            )

    def _full_record_data(self):
        return {
            "channel_name": "测试渠道",
            "start_stake": "K10+000",
            "start_stake_value": 10000.0,
            "end_stake": "K11+000",
            "end_stake_value": 11000.0,
            "section_length": 1000.0,
            "build_date": "2008-06",
            "renovation_date": None,
            "longitudinal_slope": "1/2000",
            "design_flow": 12.5,
            "channel_grade": "3级",
            "cross_section_form": "梯形",
            "embankment_top_width": 3.5,
            "inner_slope": "1:1.5",
            "outer_slope": "1:1.5",
            "increased_flow": 15.0,
            "bed_soil": "砂壤土",
            "lining_structure": "现浇混凝土衬砌",
            "freeboard": 0.6,
            "bottom_width": 4.2,
            "water_conveyance_loss": 0.08,
            "lining_material": "混凝土",
            "lining_thickness": 12.0,
            "concrete_strength": "C25",
            "channel_depth": 2.8,
            "channel_bottom_elevation": 1865.35,
        }

    def _full_inspection_results(self):
        return [
            {
                "item_code": item["item_code"],
                "category": item["category"],
                "item_name": item["item_name"],
                "grade": "A",
                "description": None,
                "remark": None,
            }
            for item in (LINED_CHANNEL_EVALUATION_ITEMS)
        ]

    def test_evaluation_configuration_has_12_items(
        self,
    ):
        self.assertEqual(
            len(LINED_CHANNEL_EVALUATION_ITEMS),
            12,
        )

        category_counts = {}

        for item in LINED_CHANNEL_EVALUATION_ITEMS:
            category_counts[item["category"]] = (
                category_counts.get(
                    item["category"],
                    0,
                )
                + 1
            )

        self.assertEqual(
            category_counts,
            {
                "水力条件": 5,
                "渠道断面": 5,
                "渠基": 2,
            },
        )

        sec_01 = next(
            item
            for item in (LINED_CHANNEL_EVALUATION_ITEMS)
            if item["item_code"] == "SEC_01"
        )

        self.assertIn(
            "b 类工程胀沉量指标",
            sec_01["standards"]["B"],
        )

        self.assertIn(
            "c 类工程胀沉量指标",
            sec_01["standards"]["C"],
        )

    def test_completion_requires_all_12_items(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        inspections = (self._full_inspection_results())[:-1]

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试渠道",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data=(self._full_record_data()),
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
            inspection_results=inspections,
            survey_date="2026-09-13",
            overall_grade="A",
            survey_comment="测试意见",
        )

        with self.assertRaisesRegex(
            ValueError,
            "全部12项",
        ):
            (database.complete_lined_channel_section_record(result["survey_record_id"]))

        record = database.get_lined_channel_section_record(result["survey_record_id"])

        self.assertEqual(
            record["record_status"],
            "draft",
        )

    def test_complete_then_edit_completed_record(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试渠道",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data=(self._full_record_data()),
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
            inspection_results=(self._full_inspection_results()),
            survey_date="2026-09-13",
            overall_grade="B",
            survey_comment="完成前意见",
        )

        complete_result = database.complete_lined_channel_section_record(
            result["survey_record_id"]
        )

        self.assertEqual(
            complete_result["inspection_count"],
            12,
        )

        completed = database.get_lined_channel_section_record(
            result["survey_record_id"]
        )

        self.assertEqual(
            completed["record_status"],
            "completed",
        )

        updated_data = self._full_record_data()

        updated_data["section_length"] = 1050.0

        (
            database.update_lined_channel_section_draft(
                survey_record_id=(result["survey_record_id"]),
                asset_name="测试渠道修改后",
                organization_unit_id=(self.office_id),
                canal_unit_id=self.canal_id,
                business_code=("1-01-01-01-001"),
                record_data=updated_data,
                start_stake_text="K10+000",
                start_stake_value=10000.0,
                end_stake_text="K11+000",
                end_stake_value=11000.0,
                inspection_results=(self._full_inspection_results()),
                survey_date="2026-09-14",
                overall_grade="C",
                survey_comment="完成后修改意见",
            )
        )

        loaded = database.get_lined_channel_section_record(result["survey_record_id"])

        self.assertEqual(
            loaded["record_status"],
            "completed",
        )

        self.assertEqual(
            loaded["asset_name"],
            "测试渠道修改后",
        )

        self.assertEqual(
            loaded["record_data"]["section_length"],
            1050.0,
        )

        self.assertEqual(
            loaded["overall_grade"],
            "C",
        )

        self.assertEqual(
            loaded["survey_comment"],
            "完成后修改意见",
        )

        inspection_results = database.get_inspection_results(result["survey_record_id"])

        self.assertEqual(
            len(inspection_results),
            12,
        )

    def test_duplicate_lined_channel_section_is_rejected(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试渠段A",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": "测试渠段A",
            },
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
        )

        with self.assertRaisesRegex(
            ValueError,
            "相同起止桩号",
        ):
            database.create_lined_channel_section_survey(
                project_id=self.project_id,
                survey_batch_id=self.batch_id,
                form_version_id=(form_version["id"]),
                asset_name="测试渠段B",
                organization_unit_id=(self.office_id),
                canal_unit_id=self.canal_id,
                business_code=("1-01-01-01-002"),
                record_data={
                    "channel_name": "测试渠段B",
                },
                start_stake_text="K10+000",
                start_stake_value=10000.0,
                end_stake_text="K11+000",
                end_stake_value=11000.0,
            )

        with database.get_connection() as connection:
            asset_count = connection.execute("""
                SELECT COUNT(*) AS count
                FROM engineering_assets
                """).fetchone()["count"]

            record_count = connection.execute("""
                SELECT COUNT(*) AS count
                FROM survey_records
                """).fetchone()["count"]

        self.assertEqual(
            asset_count,
            1,
        )

        self.assertEqual(
            record_count,
            1,
        )

    def test_update_to_duplicate_section_is_rejected(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试渠段A",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": "测试渠段A",
            },
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
        )

        second = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="测试渠段B",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-002"),
            record_data={
                "channel_name": "测试渠段B",
            },
            start_stake_text="K11+000",
            start_stake_value=11000.0,
            end_stake_text="K12+000",
            end_stake_value=12000.0,
        )

        with self.assertRaisesRegex(
            ValueError,
            "相同起止桩号",
        ):
            (
                database.update_lined_channel_section_draft(
                    survey_record_id=(second["survey_record_id"]),
                    asset_name="测试渠段B",
                    organization_unit_id=(self.office_id),
                    canal_unit_id=(self.canal_id),
                    business_code=("1-01-01-01-002"),
                    record_data={
                        "channel_name": ("测试渠段B"),
                    },
                    start_stake_text=("K10+000"),
                    start_stake_value=(10000.0),
                    end_stake_text=("K11+000"),
                    end_stake_value=(11000.0),
                )
            )

        loaded = database.get_lined_channel_section_record(second["survey_record_id"])

        self.assertEqual(
            loaded["start_stake_text"],
            "K11+000",
        )

        self.assertEqual(
            loaded["end_stake_text"],
            "K12+000",
        )

    def test_delete_record_cleans_orphan_asset(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="待删除渠段",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": ("待删除渠段"),
            },
            start_stake_text="K20+000",
            start_stake_value=20000.0,
            end_stake_text="K21+000",
            end_stake_value=21000.0,
            inspection_results=[
                {
                    "item_code": "HYD_01",
                    "category": "水力条件",
                    "item_name": "沿程流态",
                    "grade": "A",
                    "description": None,
                    "remark": None,
                }
            ],
        )

        delete_result = database.delete_lined_channel_section_record(
            result["survey_record_id"]
        )

        self.assertTrue(delete_result["asset_deleted"])

        with database.get_connection() as connection:
            record_count = connection.execute(
                """
                    SELECT COUNT(*) AS count
                    FROM survey_records
                    WHERE id = ?
                    """,
                (result["survey_record_id"],),
            ).fetchone()["count"]

            asset_count = connection.execute(
                """
                    SELECT COUNT(*) AS count
                    FROM engineering_assets
                    WHERE id = ?
                    """,
                (result["engineering_asset_id"],),
            ).fetchone()["count"]

            inspection_count = connection.execute(
                """
                    SELECT COUNT(*) AS count
                    FROM inspection_results
                    WHERE survey_record_id = ?
                    """,
                (result["survey_record_id"],),
            ).fetchone()["count"]

        self.assertEqual(
            record_count,
            0,
        )

        self.assertEqual(
            asset_count,
            0,
        )

        self.assertEqual(
            inspection_count,
            0,
        )

    def test_export_lined_channel_summary(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="汇总导出测试渠段",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data=(self._full_record_data()),
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
            inspection_results=(self._full_inspection_results()),
            survey_date="2026-09-13",
            overall_grade="B",
            survey_comment="汇总导出测试意见",
        )

        self.assertIsNotNone(result["survey_record_id"])

        records = database.get_lined_channel_section_records(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
        )

        file_path = Path(self.temp_directory.name) / "lined_channel_summary.xlsx"

        export_result = export_lined_channel_summary(
            records=records,
            file_path=file_path,
        )

        self.assertEqual(
            export_result["exported_count"],
            1,
        )

        self.assertTrue(file_path.exists())

        workbook = load_workbook(file_path)

        worksheet = workbook["渠道渠段调查汇总"]

        self.assertEqual(
            worksheet["A1"].value,
            "序号",
        )

        self.assertEqual(
            worksheet["C2"].value,
            "汇总导出测试渠段",
        )

        self.assertEqual(
            worksheet["G2"].value,
            "K10+000",
        )

        self.assertEqual(
            worksheet["H2"].value,
            "K11+000",
        )

        self.assertEqual(
            worksheet["I2"].value,
            1000.0,
        )

        # 第30列开始为12项评价
        self.assertEqual(
            worksheet.cell(
                row=2,
                column=30,
            ).value,
            "A",
        )

        # 第42列为工程状况类别
        self.assertEqual(
            worksheet.cell(
                row=2,
                column=42,
            ).value,
            "B",
        )

        workbook.close()

    def test_export_lined_channel_original_form(
        self,
    ):
        form_version = database.get_current_form_version("form_2_1")

        result = database.create_lined_channel_section_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(form_version["id"]),
            asset_name="原表导出测试渠段",
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=("1-01-01-01-001"),
            record_data=(self._full_record_data()),
            start_stake_text="K10+000",
            start_stake_value=10000.0,
            end_stake_text="K11+000",
            end_stake_value=11000.0,
            inspection_results=(self._full_inspection_results()),
            survey_date="2026-09-13",
            overall_grade="C",
            survey_comment="原表导出测试意见",
        )

        file_path = Path(self.temp_directory.name) / "lined_channel_original.xlsx"

        export_result = export_lined_channel_original_form(
            survey_record_id=(result["survey_record_id"]),
            file_path=file_path,
        )

        self.assertTrue(file_path.exists())

        self.assertEqual(
            export_result["business_code"],
            "1-01-01-01-001",
        )

        workbook = load_workbook(file_path)

        worksheet = workbook["附表2.1"]

        self.assertEqual(
            worksheet["B5"].value,
            "原表导出测试渠段",
        )

        self.assertEqual(
            worksheet["H5"].value,
            "K10+000 ～ K11+000",
        )

        self.assertEqual(
            worksheet["B6"].value,
            1000.0,
        )

        self.assertEqual(
            worksheet["D6"].value,
            "2008-06",
        )

        self.assertEqual(
            worksheet["H6"].value,
            "1/2000",
        )

        self.assertEqual(
            worksheet["E11"].value,
            "A",
        )

        self.assertEqual(
            worksheet["E22"].value,
            "A",
        )

        self.assertEqual(
            worksheet["C23"].value,
            "原表导出测试意见",
        )

        self.assertEqual(
            worksheet["J23"].value,
            "C",
        )

        self.assertEqual(
            worksheet["J24"].value,
            "2026-09-13",
        )

        self.assertEqual(
            worksheet["H7"].value,
            "1:1.5/1:1.5",
        )

        workbook.close()


def test_lined_channel_update_can_change_ownership(
    self,
):
    form_version = database.get_current_form_version("form_2_1")

    result = database.create_lined_channel_section_survey(
        project_id=self.project_id,
        survey_batch_id=self.batch_id,
        form_version_id=(form_version["id"]),
        asset_name="原渠段",
        organization_unit_id=(self.office_id),
        canal_unit_id=(self.canal_id),
        business_code=("1-01-01-01-001"),
        record_data={
            "channel_name": "原渠段",
        },
        start_stake_text=("K10+000"),
        start_stake_value=(10000.0),
        end_stake_text=("K11+000"),
        end_stake_value=(11000.0),
    )

    new_office_id = database.create_organization_unit(
        name="测试第二水管所",
        unit_type="water_office",
        business_code="02",
        parent_id=(self.department_id),
    )

    self.assertIsNotNone(new_office_id)

    assert new_office_id is not None

    new_canal_id = database.create_canal_unit(
        name="测试第二干渠",
        canal_level="01",
        parent_id=None,
        organization_unit_id=(new_office_id),
        description=None,
    )

    self.assertIsNotNone(new_canal_id)

    assert new_canal_id is not None

    database.update_lined_channel_section_draft(
        survey_record_id=(result["survey_record_id"]),
        asset_name="修改归属后的渠段",
        organization_unit_id=(new_office_id),
        canal_unit_id=(new_canal_id),
        business_code=("1-02-01-01-001"),
        record_data={
            "channel_name": "修改归属后的渠段",
        },
        start_stake_text=("K10+000"),
        start_stake_value=(10000.0),
        end_stake_text=("K11+000"),
        end_stake_value=(11000.0),
    )

    loaded = database.get_lined_channel_section_record(result["survey_record_id"])

    self.assertEqual(
        loaded["organization_unit_id"],
        new_office_id,
    )

    self.assertEqual(
        loaded["office_id"],
        new_office_id,
    )

    self.assertEqual(
        loaded["canal_unit_id"],
        new_canal_id,
    )

    self.assertEqual(
        loaded["canal_id"],
        new_canal_id,
    )

    self.assertEqual(
        loaded["business_code"],
        "1-02-01-01-001",
    )

    # EngineeringAsset 与
    # SurveyRecord 必须同步更新。
    with database.get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                ea.organization_unit_id
                    AS asset_office_id,
                ea.canal_unit_id
                    AS asset_canal_id,
                ea.business_code
                    AS asset_business_code,

                sr.organization_unit_id
                    AS record_office_id,
                sr.canal_unit_id
                    AS record_canal_id,
                sr.business_code
                    AS record_business_code

            FROM survey_records AS sr

            JOIN engineering_assets AS ea
                ON sr.engineering_asset_id
                    = ea.id

            WHERE sr.id = ?
            """,
            (result["survey_record_id"],),
        ).fetchone()

    self.assertIsNotNone(rows)

    assert rows is not None

    self.assertEqual(
        rows["asset_office_id"],
        new_office_id,
    )

    self.assertEqual(
        rows["record_office_id"],
        new_office_id,
    )

    self.assertEqual(
        rows["asset_canal_id"],
        new_canal_id,
    )

    self.assertEqual(
        rows["record_canal_id"],
        new_canal_id,
    )

    self.assertEqual(
        rows["asset_business_code"],
        "1-02-01-01-001",
    )

    self.assertEqual(
        rows["record_business_code"],
        "1-02-01-01-001",
    )


if __name__ == "__main__":
    unittest.main()
