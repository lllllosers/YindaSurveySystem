import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from forms.engineering.form_2_10 import (
    FORM_2_10,
)
from forms.engineering.form_2_11 import (
    FORM_2_11,
)
from forms.engineering.form_2_14 import (
    FORM_2_14,
)

from forms.engineering.persistence import (
    complete_engineering_record,
    create_engineering_record,
    get_engineering_record,
    update_engineering_record,
)

from services.engineering_original_form_export import (
    export_engineering_original_form,
)
from services.engineering_summary_export import (
    export_engineering_summary,
)

from tests.workflow_test_support import (
    EngineeringWorkflowTestCaseBase,
)


class BatchBWorkflowMixin:
    """
    Batch B 三张 A/B/C 三级评价工程调查表
    的生产级 workflow 合同。

    每张表验证：
    - 新增草稿与 point 工程位置；
    - 同位置重复保护；
    - A/B/C 完成调查；
    - completed 后继续修改；
    - D 级明确拒绝；
    - DataQuery 共用查询结果；
    - 详细汇总实际导出；
    - 正式原表实际导出。
    """

    DEFINITION = None
    ASSET_NAME = None

    POINT_STAKE = "K80+500"
    POINT_STAKE_VALUE = 80500.0

    def setUp(self):
        super().setUp()

        if self.DEFINITION is None:
            raise ValueError(
                "Batch B workflow 必须配置 DEFINITION。"
            )

        self.assertEqual(
            self.DEFINITION.form_code,
            self.FORM_CODE,
        )
        self.assertEqual(
            self.DEFINITION.position.kind,
            "point",
        )
        self.assertEqual(
            self.DEFINITION.grade_options,
            (
                "A",
                "B",
                "C",
            ),
        )

    # =========================================================
    # payload helpers
    # =========================================================

    def _record_data(
        self,
        *,
        asset_name=None,
    ):
        asset_name = (
            asset_name
            or self.ASSET_NAME
        )

        result = {}

        for index, field in enumerate(
            self.DEFINITION.fields,
            start=1,
        ):
            if (
                field.key
                == self.DEFINITION.asset_name_field
            ):
                value = asset_name

            elif field.input_type == "stake":
                value = self.POINT_STAKE

            elif field.input_type == "month":
                value = (
                    "2024-01"
                    if not field.required
                    else "2020-06"
                )

            elif field.input_type == "choice":
                self.assertTrue(
                    field.choices
                )
                value = field.choices[0]

            elif field.input_type == "integer":
                value = (
                    index + 1
                )

            elif (
                field.input_type
                == "signed_decimal"
            ):
                value = -(
                    1000.0
                    + index
                    + 0.25
                )

            elif field.input_type == "decimal":
                value = (
                    10.0
                    + index
                    + 0.25
                )

            elif field.input_type == "structure_grade":
                value = "3级"

            elif field.input_type == "concrete_strength":
                value = "C30"

            elif field.input_type == "text":
                value = (
                    f"测试{field.label}"
                )

            else:
                raise AssertionError(
                    "测试尚未覆盖字段类型："
                    f"{field.input_type}"
                )

            result[field.key] = value

        return result

    def _inspection_results(
        self,
        grade="B",
    ):
        return [
            {
                "item_code": (
                    item["item_code"]
                ),
                "category": (
                    item["category"]
                ),
                "item_name": (
                    item["item_name"]
                ),
                "grade": grade,
                "description": None,
                "remark": None,
            }
            for item
            in self.DEFINITION.evaluation_items
        ]

    def _payload(
        self,
        *,
        asset_name=None,
        grade="B",
        overall_grade="B",
        survey_comment=(
            "Batch B workflow 测试。"
        ),
    ):
        asset_name = (
            asset_name
            or self.ASSET_NAME
        )

        return {
            "asset_name": asset_name,
            "record_data": (
                self._record_data(
                    asset_name=asset_name
                )
            ),
            "position": {
                "kind": "point",
                "single_stake_text": (
                    self.POINT_STAKE
                ),
                "single_stake_value": (
                    self.POINT_STAKE_VALUE
                ),
            },
            "inspection_results": (
                self._inspection_results(
                    grade=grade
                )
            ),
            "survey_date": "2026-09-17",
            "overall_grade": overall_grade,
            "survey_comment": survey_comment,
        }

    def _business_code(
        self,
        sequence=1,
    ):
        return (
            "1-01-01-"
            f"{self.DEFINITION.business_type_code}"
            f"-{sequence:03d}"
        )

    def _create_draft(
        self,
        *,
        sequence=1,
        asset_name=None,
        payload=None,
    ):
        if payload is None:
            payload = self._payload(
                asset_name=asset_name
            )

        return create_engineering_record(
            self.DEFINITION,
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=int(
                self.form_version["id"]
            ),
            organization_unit_id=(
                self.office_id
            ),
            canal_unit_id=self.canal_id,
            business_code=(
                self._business_code(
                    sequence
                )
            ),
            payload=payload,
        )

    @staticmethod
    def _resolve_original_binding(
        *,
        record,
        binding,
    ):
        record_data = (
            record["record_data"]
            or {}
        )

        if binding.source == "record":
            values = [
                record[key]
                for key in binding.keys
            ]

        elif (
            binding.source
            == "record_data"
        ):
            values = [
                record_data.get(key)
                for key in binding.keys
            ]

        else:
            raise AssertionError(
                "正式原表测试出现未支持数据源："
                f"{binding.source}"
            )

        if binding.formatter is not None:
            return binding.formatter(
                *values
            )

        return values[0]

    # =========================================================
    # 1. 草稿 / point / 查询 / 重复保护
    # =========================================================

    def test_draft_query_and_duplicate_protection(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        record = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(
            record
        )
        assert record is not None

        self.assertEqual(
            record["record_status"],
            "draft",
        )
        self.assertEqual(
            record["asset_type"],
            self.DEFINITION.asset_type,
        )
        self.assertEqual(
            record["asset_name"],
            self.ASSET_NAME,
        )
        self.assertEqual(
            record[
                "single_stake_text"
            ],
            self.POINT_STAKE,
        )
        self.assertAlmostEqual(
            float(
                record[
                    "single_stake_value"
                ]
            ),
            self.POINT_STAKE_VALUE,
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=(
                    self.project_id
                ),
                survey_batch_id=(
                    self.batch_id
                ),
                form_code=(
                    self.DEFINITION
                    .form_code
                ),
            )
        )

        self.assertEqual(
            len(records),
            1,
        )
        self.assertEqual(
            records[0][
                "record_status"
            ],
            "draft",
        )
        self.assertIn(
            self.POINT_STAKE,
            (
                records[0][
                    "engineering_position"
                ]
                or ""
            ),
        )

        with self.assertRaises(
            ValueError
        ):
            self._create_draft(
                sequence=2,
                asset_name=(
                    self.ASSET_NAME
                    + "重复"
                ),
            )

    # =========================================================
    # 2. 完成 -> completed 修改
    # =========================================================

    def test_complete_then_edit_completed_record(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        completion = (
            complete_engineering_record(
                self.DEFINITION,
                survey_record_id=(
                    survey_record_id
                ),
                payload=self._payload(),
            )
        )

        self.assertEqual(
            completion[
                "inspection_count"
            ],
            len(
                self.DEFINITION
                .evaluation_items
            ),
        )

        completed = (
            get_engineering_record(
                self.DEFINITION,
                survey_record_id=(
                    survey_record_id
                ),
            )
        )

        self.assertIsNotNone(
            completed
        )
        assert completed is not None

        self.assertEqual(
            completed[
                "record_status"
            ],
            "completed",
        )
        self.assertEqual(
            completed[
                "overall_grade"
            ],
            "B",
        )

        changed_name = (
            self.ASSET_NAME
            + "（已完成修改）"
        )

        update_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(
                asset_name=changed_name,
                grade="C",
                overall_grade="C",
                survey_comment=(
                    "completed 后三级评价修改成功。"
                ),
            ),
        )

        updated = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(
            updated
        )
        assert updated is not None

        self.assertEqual(
            updated[
                "record_status"
            ],
            "completed",
        )
        self.assertEqual(
            updated["asset_name"],
            changed_name,
        )
        self.assertEqual(
            updated[
                "overall_grade"
            ],
            "C",
        )
        self.assertEqual(
            updated[
                "survey_comment"
            ],
            "completed 后三级评价修改成功。",
        )

        inspections = (
            database
            .get_inspection_results(
                survey_record_id
            )
        )

        self.assertTrue(
            inspections
        )
        self.assertTrue(
            all(
                row["grade"] == "C"
                for row in inspections
            )
        )

    # =========================================================
    # 3. D 级必须被三级评价合同拒绝
    # =========================================================

    def test_grade_d_is_rejected(
        self,
    ):
        payload = self._payload(
            grade="D",
            overall_grade="D",
        )

        result = self._create_draft(
            payload=payload,
        )

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        with self.assertRaisesRegex(
            ValueError,
            "评价等级无效|工程状况类别无效",
        ):
            complete_engineering_record(
                self.DEFINITION,
                survey_record_id=(
                    survey_record_id
                ),
                payload=payload,
            )

        record = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(
            record
        )
        assert record is not None

        self.assertEqual(
            record[
                "record_status"
            ],
            "draft",
        )

    # =========================================================
    # 4. 详细汇总真实导出
    # =========================================================

    def test_summary_export_real_workbook(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        complete_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(),
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=(
                    self.project_id
                ),
                survey_batch_id=(
                    self.batch_id
                ),
                form_code=(
                    self.DEFINITION
                    .form_code
                ),
            )
        )

        output = (
            Path(
                self.temp_directory.name
            )
            / (
                self.DEFINITION
                .form_code
                + "_summary.xlsx"
            )
        )

        export_result = (
            export_engineering_summary(
                self.DEFINITION,
                records=records,
                file_path=output,
            )
        )

        self.assertEqual(
            export_result[
                "exported_count"
            ],
            1,
        )
        self.assertTrue(
            output.exists()
        )

        workbook = load_workbook(
            output,
            data_only=False,
        )

        try:
            summary = (
                self.DEFINITION
                .summary_export_definition
            )

            self.assertIsNotNone(
                summary
            )
            assert summary is not None

            worksheet = workbook[
                summary.sheet_name
            ]

            headers = {
                cell.value: cell.column
                for cell
                in worksheet[1]
            }

            for header in (
                "业务编号",
                "名称",
                "工程状况类别",
                "状态",
            ):
                self.assertIn(
                    header,
                    headers,
                )

            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "业务编号"
                    ],
                ).value,
                self._business_code(),
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "名称"
                    ],
                ).value,
                self.ASSET_NAME,
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "工程状况类别"
                    ],
                ).value,
                "B",
            )
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=headers[
                        "状态"
                    ],
                ).value,
                "录入完成",
            )

            for item in (
                self.DEFINITION
                .evaluation_items
            ):
                header = (
                    f"{item['category']}"
                    f"-{item['item_name']}"
                )

                self.assertIn(
                    header,
                    headers,
                )
                self.assertEqual(
                    worksheet.cell(
                        row=2,
                        column=headers[
                            header
                        ],
                    ).value,
                    "B",
                )

        finally:
            workbook.close()

    # =========================================================
    # 5. 正式原表真实导出
    # =========================================================

    def test_original_form_real_export(
        self,
    ):
        result = self._create_draft()

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        complete_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            payload=self._payload(),
        )

        original = (
            self.DEFINITION
            .original_form_export_definition
        )

        self.assertIsNotNone(
            original
        )
        assert original is not None

        output = (
            Path(
                self.temp_directory.name
            )
            / (
                self.DEFINITION
                .form_code
                + "_original.xlsx"
            )
        )

        export_engineering_original_form(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
            file_path=output,
        )

        self.assertTrue(
            output.exists()
        )

        record = get_engineering_record(
            self.DEFINITION,
            survey_record_id=(
                survey_record_id
            ),
        )

        self.assertIsNotNone(
            record
        )
        assert record is not None

        workbook = load_workbook(
            output,
            data_only=False,
        )

        try:
            worksheet = workbook[
                original.sheet_name
            ]

            self.assertEqual(
                worksheet["J3"].value,
                self._business_code(),
            )

            for cell_binding in (
                original.field_bindings
            ):
                expected = (
                    self
                    ._resolve_original_binding(
                        record=record,
                        binding=(
                            cell_binding
                            .binding
                        ),
                    )
                )

                actual = worksheet[
                    cell_binding.cell
                ].value

                self.assertEqual(
                    actual,
                    expected,
                    msg=(
                        self.DEFINITION
                        .form_code
                        + " "
                        + cell_binding.cell
                    ),
                )

            start_row = (
                original
                .evaluation_binding
                .start_row
            )
            grade_column = (
                original
                .evaluation_binding
                .column
            )

            for offset in range(
                len(
                    self.DEFINITION
                    .evaluation_items
                )
            ):
                self.assertEqual(
                    worksheet[
                        f"{grade_column}"
                        f"{start_row + offset}"
                    ].value,
                    "B",
                )

            conclusion = (
                original
                .conclusion_binding
            )

            self.assertEqual(
                worksheet[
                    conclusion
                    .survey_comment_cell
                ].value,
                "Batch B workflow 测试。",
            )
            self.assertEqual(
                worksheet[
                    conclusion
                    .overall_grade_cell
                ].value,
                "B",
            )
            self.assertEqual(
                worksheet[
                    conclusion
                    .survey_date_cell
                ].value,
                "2026-09-17",
            )

            # 三级评价正式表情况描述不得出现 d 级。
            for offset in range(
                len(
                    self.DEFINITION
                    .evaluation_items
                )
            ):
                row_number = (
                    start_row + offset
                )

                description = (
                    worksheet[
                        f"F{row_number}"
                    ].value
                    or ""
                )

                self.assertIn(
                    "a ",
                    description,
                )
                self.assertIn(
                    "b ",
                    description,
                )
                self.assertIn(
                    "c ",
                    description,
                )
                self.assertNotIn(
                    "d ",
                    description,
                )

        finally:
            workbook.close()


class Form210WorkflowTestCase(
    BatchBWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_10"
    DEFINITION = FORM_2_10
    TEST_DB_FILENAME = (
        "test_form_2_10_batch_b.db"
    )
    ASSET_NAME = (
        "测试标准断面量水设施"
    )


class Form211WorkflowTestCase(
    BatchBWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_11"
    DEFINITION = FORM_2_11
    TEST_DB_FILENAME = (
        "test_form_2_11_batch_b.db"
    )
    ASSET_NAME = (
        "测试堰槽量水设施"
    )


class Form214WorkflowTestCase(
    BatchBWorkflowMixin,
    EngineeringWorkflowTestCaseBase,
):
    FORM_CODE = "form_2_14"
    DEFINITION = FORM_2_14
    TEST_DB_FILENAME = (
        "test_form_2_14_batch_b.db"
    )
    ASSET_NAME = "测试沟段"


if __name__ == "__main__":
    unittest.main()
