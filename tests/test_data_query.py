import gc
import sys
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from forms.engineering.form_2_1 import (
    FORM_2_1,
)

from services.query_export import (
    export_common_query_summary,
)
from services.aqueduct_export import (
    export_aqueduct_summary,
)
from services.aqueduct_evaluation import (
    AQUEDUCT_EVALUATION_ITEMS,
)
from pages.data_query_page import (
    ENGINEERING_SUMMARY_EXPORTERS,
)
from services.culvert_export import (
    export_culvert_summary,
)


def create_form_2_1_query_fixture(
    project_id,
    survey_batch_id,
    form_version_id,
    asset_name,
    organization_unit_id,
    canal_unit_id,
    business_code,
    record_data,
    start_stake_text=None,
    start_stake_value=None,
    end_stake_text=None,
    end_stake_value=None,
    inspection_results=None,
    survey_date=None,
    overall_grade=None,
    survey_comment=None,
):
    """
    DataQuery 测试只负责准备
    一条附表2.1 range 工程记录。

    使用公共 range DB API，
    不再依赖附表2.1 legacy wrapper。
    """

    return (
        database
        .create_range_engineering_survey(
            project_id=project_id,
            survey_batch_id=(
                survey_batch_id
            ),
            form_version_id=(
                form_version_id
            ),
            asset_name=asset_name,
            asset_type=(
                FORM_2_1.asset_type
            ),
            organization_unit_id=(
                organization_unit_id
            ),
            canal_unit_id=(
                canal_unit_id
            ),
            business_code=(
                business_code
            ),
            record_data=record_data,
            start_stake_text=(
                start_stake_text
            ),
            start_stake_value=(
                start_stake_value
            ),
            end_stake_text=(
                end_stake_text
            ),
            end_stake_value=(
                end_stake_value
            ),
            inspection_results=(
                inspection_results
            ),
            survey_date=survey_date,
            overall_grade=(
                overall_grade
            ),
            survey_comment=(
                survey_comment
            ),
        )
    )


class DataQueryTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()

        self.temp_data_dir = Path(self.temp_directory.name) / "local_data"

        self.temp_db_path = self.temp_data_dir / "test_data_query.db"

        self.original_data_dir = database.DATA_DIR

        self.original_db_path = database.DB_PATH

        database.DATA_DIR = self.temp_data_dir

        database.DB_PATH = self.temp_db_path

        database.init_database()
        database.create_initial_forms()

        # =====================================================
        # 项目
        # =====================================================

        project_result = database.create_project(
            name="统一数据查询测试项目",
            short_name="查询测试",
        )

        self.project_id = int(project_result["project_id"])

        # =====================================================
        # 两个调查批次
        # =====================================================

        batch_1 = database.create_survey_batch(
            project_id=(self.project_id),
            batch_name=("2026年度调查"),
            batch_code=("QUERY_2026"),
            start_date=("2026-09-01"),
            end_date=("2026-12-31"),
        )

        self.batch_1_id = int(batch_1["batch_id"])

        batch_2 = database.create_survey_batch(
            project_id=(self.project_id),
            batch_name=("2027年度调查"),
            batch_code=("QUERY_2027"),
            start_date=("2027-09-01"),
            end_date=("2027-12-31"),
        )

        self.batch_2_id = int(batch_2["batch_id"])

        # =====================================================
        # 机构 / 渠系
        # =====================================================

        department_id = database.create_organization_unit(
            name="测试基层处",
            unit_type="department",
            business_code="1",
        )

        self.office_id = database.create_organization_unit(
            name="测试水管所",
            unit_type="water_office",
            business_code="01",
            parent_id=department_id,
        )

        self.canal_id = database.create_canal_unit(
            name="测试总干渠",
            canal_level="01",
            organization_unit_id=(self.office_id),
        )

        # =====================================================
        # 附表2.1：2026
        # =====================================================

        form_2_1 = database.get_current_form_version("form_2_1")

        create_form_2_1_query_fixture(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
            form_version_id=(form_2_1["id"]),
            asset_name="测试渠段A",
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=("1-01-01-01-001"),
            record_data={
                "channel_name": ("测试渠段A"),
                "section_length": (1000.0),
            },
            start_stake_text=("K10+000"),
            start_stake_value=(10000.0),
            end_stake_text=("K11+000"),
            end_stake_value=(11000.0),
            survey_date=("2026-09-13"),
            overall_grade="B",
            survey_comment=("测试意见A"),
        )

        # =====================================================
        # 附表2.2：2026
        # =====================================================

        form_2_2 = database.get_current_form_version("form_2_2")

        database.create_engineering_survey(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
            form_version_id=(form_2_2["id"]),
            asset_name="测试水闸A",
            asset_type="sluice_gate",
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=("1-01-01-02-001"),
            record_data={
                "stake": "K20+500",
                "design_flow": 10.0,
            },
            single_stake_text=("K20+500"),
            single_stake_value=(20500.0),
            survey_date=("2026-09-14"),
            overall_grade="C",
            survey_comment=("测试意见B"),
        )

        # =====================================================
        # 附表2.3：2026
        # =====================================================

        form_2_3 = database.get_current_form_version("form_2_3")

        database.create_engineering_survey(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
            form_version_id=(form_2_3["id"]),
            asset_name="测试渡槽A",
            asset_type="aqueduct",
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=("1-01-01-03-001"),
            record_data={
                "stake": "K25+300",
                "design_flow": 8.0,
                "length": 120.0,
            },
            single_stake_text=("K25+300"),
            single_stake_value=(25300.0),
            survey_date=("2026-09-15"),
            overall_grade="D",
            survey_comment=("测试意见D"),
        )

        # =====================================================
        # 附表2.3：2026
        # 详细汇总导出测试记录
        # =====================================================

        aqueduct_record_data = {
            "stake": "K26+100",
            "stake_value": 26100.0,
            "design_flow": 8.5,
            "structure_grade": "3级",
            "build_date": "2010-06",
            "renovation_date": "2021-09",
            "length": 150.0,
            "increased_flow": 10.0,
            "structure_form": "梁式渡槽",
            "section_width": 3.2,
            "section_height": 2.4,
            "trough_body_structure": ("钢筋混凝土"),
            "trough_wall_thickness": 0.25,
            "waterstop_form": "橡胶止水",
            "trough_bottom_elevation": -3.5,
            "span_count": 5,
            "lower_support_structure_form": ("排架式"),
        }

        aqueduct_inspections = [
            {
                "item_code": (item["item_code"]),
                "category": (item["category"]),
                "item_name": (item["item_name"]),
                "grade": "B",
                "description": None,
                "remark": None,
            }
            for item in AQUEDUCT_EVALUATION_ITEMS
        ]

        database.create_engineering_survey(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
            form_version_id=(form_2_3["id"]),
            asset_name=("详细汇总测试渡槽"),
            asset_type="aqueduct",
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=("1-01-01-03-002"),
            record_data=(aqueduct_record_data),
            single_stake_text="K26+100",
            single_stake_value=26100.0,
            inspection_results=(aqueduct_inspections),
            survey_date="2026-09-16",
            overall_grade="B",
            survey_comment=("详细汇总导出测试。"),
        )

        # =====================================================
        # 附表2.1：2027
        # =====================================================

        create_form_2_1_query_fixture(
            project_id=self.project_id,
            survey_batch_id=(self.batch_2_id),
            form_version_id=(form_2_1["id"]),
            asset_name="测试渠段B",
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=("1-01-01-01-002"),
            record_data={
                "channel_name": ("测试渠段B"),
                "section_length": (500.0),
            },
            start_stake_text=("K30+000"),
            start_stake_value=(30000.0),
            end_stake_text=("K30+500"),
            end_stake_value=(30500.0),
            survey_date=("2027-09-13"),
            overall_grade="A",
            survey_comment=("测试意见C"),
        )

    def tearDown(self):
        database.DATA_DIR = self.original_data_dir

        database.DB_PATH = self.original_db_path

        gc.collect()
        time.sleep(0.05)

        self.temp_directory.cleanup()

    # =========================================================
    # 测试1：3个表统一查询
    # =========================================================

    def test_query_returns_multiple_forms(
        self,
    ):
        records = database.get_engineering_survey_query_records(
            project_id=self.project_id
        )

        self.assertEqual(
            len(records),
            5,
        )

        form_codes = {record["form_code"] for record in records}

        self.assertEqual(
            form_codes,
            {
                "form_2_1",
                "form_2_2",
                "form_2_3",
            },
        )

        lined_record = next(
            record for record in records if record["asset_name"] == "测试渠段A"
        )

        self.assertEqual(
            lined_record["engineering_position"],
            "K10+000 ～ K11+000",
        )

        sluice_record = next(
            record for record in records if record["asset_name"] == "测试水闸A"
        )

        self.assertEqual(
            sluice_record["engineering_position"],
            "K20+500",
        )

        aqueduct_record = next(
            record for record in records if record["asset_name"] == "测试渡槽A"
        )

        self.assertEqual(
            aqueduct_record["engineering_position"],
            "K25+300",
        )

        self.assertEqual(
            aqueduct_record["form_code"],
            "form_2_3",
        )

        self.assertEqual(
            aqueduct_record["overall_grade"],
            "D",
        )

    # =========================================================
    # 测试2：数据库层按批次 / 表单过滤
    # =========================================================

    def test_query_filters_batch_and_form(
        self,
    ):
        batch_records = database.get_engineering_survey_query_records(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
        )

        self.assertEqual(
            len(batch_records),
            4,
        )

        form_records = database.get_engineering_survey_query_records(
            project_id=self.project_id,
            form_code="form_2_1",
        )

        self.assertEqual(
            len(form_records),
            2,
        )

        sluice_records = database.get_engineering_survey_query_records(
            project_id=self.project_id,
            survey_batch_id=(self.batch_1_id),
            form_code="form_2_2",
        )

        self.assertEqual(
            len(sluice_records),
            1,
        )

        self.assertEqual(
            sluice_records[0]["asset_name"],
            "测试水闸A",
        )

        aqueduct_records = database.get_engineering_survey_query_records(
            project_id=(self.project_id),
            survey_batch_id=(self.batch_1_id),
            form_code="form_2_3",
        )

        self.assertEqual(
            len(aqueduct_records),
            2,
        )

        aqueduct_record = next(
            record for record in aqueduct_records if record["asset_name"] == "测试渡槽A"
        )

        self.assertEqual(
            aqueduct_record["engineering_position"],
            "K25+300",
        )

        detailed_record = next(
            record
            for record in aqueduct_records
            if record["asset_name"] == "详细汇总测试渡槽"
        )

        self.assertEqual(
            detailed_record["engineering_position"],
            "K26+100",
        )

    # =========================================================
    # 测试3：跨表公共结果导出
    # =========================================================

    def test_export_common_query_summary(
        self,
    ):
        records = database.get_engineering_survey_query_records(
            project_id=self.project_id
        )

        file_path = Path(self.temp_directory.name) / "query_summary.xlsx"

        result = export_common_query_summary(
            records=records,
            file_path=file_path,
        )

        self.assertEqual(
            result["exported_count"],
            5,
        )

        self.assertTrue(file_path.exists())

        workbook = load_workbook(file_path)

        worksheet = workbook["统一数据查询"]

        self.assertEqual(
            worksheet["A1"].value,
            "序号",
        )

        self.assertEqual(
            worksheet["B1"].value,
            "调查表",
        )

        self.assertEqual(
            worksheet["I1"].value,
            "工程位置",
        )

        exported_names = {
            worksheet.cell(
                row=row,
                column=5,
            ).value
            for row in range(
                2,
                worksheet.max_row + 1,
            )
        }

        self.assertEqual(
            exported_names,
            {
                "测试渠段A",
                "测试水闸A",
                "测试渡槽A",
                "测试渠段B",
                "详细汇总测试渡槽",
            },
        )

        # =============================================
        # 专门检查附表2.3公共字段导出
        # =============================================

        aqueduct_row = None

        for row in range(
            2,
            worksheet.max_row + 1,
        ):
            if (
                worksheet.cell(
                    row=row,
                    column=5,
                ).value
                == "测试渡槽A"
            ):
                aqueduct_row = row
                break

        self.assertIsNotNone(aqueduct_row)

        assert aqueduct_row is not None

        # I列：工程位置
        self.assertEqual(
            worksheet.cell(
                row=aqueduct_row,
                column=9,
            ).value,
            "K25+300",
        )

        # J列：工程状况类别
        self.assertEqual(
            worksheet.cell(
                row=aqueduct_row,
                column=10,
            ).value,
            "D",
        )

        # K列：调查时间
        self.assertEqual(
            worksheet.cell(
                row=aqueduct_row,
                column=11,
            ).value,
            "2026-09-15",
        )

        workbook.close()

    # =========================================================
    # 测试4：附表2.3详细汇总导出
    # =========================================================

    def test_export_aqueduct_summary(
        self,
    ):
        records = database.get_engineering_survey_query_records(
            project_id=(self.project_id),
            survey_batch_id=(self.batch_1_id),
            form_code="form_2_3",
        )

        file_path = Path(self.temp_directory.name) / "aqueduct_summary.xlsx"

        result = export_aqueduct_summary(
            records=records,
            file_path=file_path,
        )

        self.assertEqual(
            result["exported_count"],
            2,
        )

        self.assertTrue(file_path.exists())

        workbook = load_workbook(file_path)

        worksheet = workbook["渡槽调查汇总"]

        self.assertEqual(
            worksheet["A1"].value,
            "序号",
        )

        self.assertEqual(
            worksheet["C1"].value,
            "名称",
        )

        self.assertEqual(
            worksheet["G1"].value,
            "桩号",
        )

        self.assertEqual(
            worksheet["H1"].value,
            "设计流量（m³/s）",
        )

        self.assertEqual(
            worksheet["L1"].value,
            "长度",
        )

        self.assertEqual(
            worksheet["O1"].value,
            "断面尺寸（宽×高）",
        )

        self.assertEqual(
            worksheet["Q1"].value,
            "槽壁厚度",
        )

        self.assertEqual(
            worksheet["S1"].value,
            "槽底高程",
        )

        detailed_row = None

        for row in range(
            2,
            worksheet.max_row + 1,
        ):
            if (
                worksheet.cell(
                    row=row,
                    column=3,
                ).value
                == "详细汇总测试渡槽"
            ):
                detailed_row = row
                break

        self.assertIsNotNone(detailed_row)

        assert detailed_row is not None

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=7,
            ).value,
            "K26+100",
        )

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=15,
            ).value,
            "3.2×2.4",
        )

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=19,
            ).value,
            -3.5,
        )

        # 12项评价从第22列开始，
        # 本测试统一填写为B。
        for column in range(
            22,
            34,
        ):
            self.assertEqual(
                worksheet.cell(
                    row=detailed_row,
                    column=column,
                ).value,
                "B",
            )

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=34,
            ).value,
            "B",
        )

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=35,
            ).value,
            "2026-09-16",
        )

        self.assertEqual(
            worksheet.cell(
                row=detailed_row,
                column=36,
            ).value,
            "详细汇总导出测试。",
        )

        workbook.close()

    def test_form_2_6_uses_detailed_summary_exporter(
        self,
    ):
        """
        DataQuery单独选择附表2.6时，
        应使用2.6详细汇总导出器。
        """

        self.assertIn(
            "form_2_6",
            ENGINEERING_SUMMARY_EXPORTERS,
        )

        self.assertIs(
            ENGINEERING_SUMMARY_EXPORTERS[
                "form_2_6"
            ],
            export_culvert_summary,
        )

if __name__ == "__main__":
    unittest.main()
