import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


PROJECT_ROOT = Path(
    __file__
).resolve().parents[2]

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


from forms.engineering.form_2_1 import (
    FORM_2_1,
)
from services.engineering_batch_export import (
    BatchExportPlan,
    BatchExportRequest,
    FormExportGroup,
    build_batch_export_plan,
    execute_batch_export,
    load_scope_records,
    sanitize_filename,
)
from services.survey_scope import (
    SurveyScope,
)


def _record(
    *,
    survey_record_id,
    form_code="form_2_1",
    organization_unit_id=2,
    canal_unit_id=11,
    status="completed",
    canal_name="测试支渠",
):
    return {
        "survey_record_id": survey_record_id,
        "engineering_asset_id": (
            survey_record_id + 100
        ),
        "project_id": 1,
        "survey_batch_id": 2,
        "organization_unit_id": (
            organization_unit_id
        ),
        "canal_unit_id": canal_unit_id,
        "batch_name": "2026年度调查",
        "batch_code": "2026",
        "form_code": form_code,
        "form_number": (
            "2.1"
            if form_code == "form_2_1"
            else "2.2"
        ),
        "form_name": "测试调查表",
        "form_display_name": "测试调查表",
        "asset_type": "test",
        "business_code": (
            f"1-01-03-01-"
            f"{survey_record_id:03d}"
        ),
        "asset_name": (
            f"测试工程"
            f"{survey_record_id}"
        ),
        "department_name": "测试处",
        "office_name": "测试所",
        "canal_name": canal_name,
        "single_stake_text": (
            f"K{survey_record_id}+000"
        ),
        "start_stake_text": "",
        "end_stake_text": "",
        "engineering_position": (
            f"K{survey_record_id}+000"
        ),
        "record_data": {},
        "overall_grade": "A",
        "survey_date": "2026-09-17",
        "record_status": status,
        "updated_at": "2026-09-17",
    }


class EngineeringBatchExportTestCase(
    unittest.TestCase,
):
    def test_sanitize_filename_for_windows(
        self,
    ):
        self.assertEqual(
            sanitize_filename(
                'A/B:C*D?"E'
            ),
            "A_B_C_D__E",
        )

        self.assertEqual(
            sanitize_filename("CON"),
            "_CON",
        )

    @patch(
        "services.engineering_batch_export."
        "get_engineering_survey_query_records"
    )
    @patch(
        "services.engineering_batch_export."
        "get_canal_units"
    )
    @patch(
        "services.engineering_batch_export."
        "get_water_offices"
    )
    @patch(
        "services.engineering_batch_export."
        "get_departments"
    )
    def test_scope_loader_applies_descendants_and_completed(
        self,
        mock_departments,
        mock_offices,
        mock_canals,
        mock_records,
    ):
        mock_departments.return_value = [
            {
                "id": 1,
                "parent_id": None,
            }
        ]

        mock_offices.return_value = [
            {
                "id": 2,
                "parent_id": 1,
            },
            {
                "id": 3,
                "parent_id": 1,
            },
        ]

        mock_canals.return_value = [
            {
                "id": 10,
                "parent_id": None,
            },
            {
                "id": 11,
                "parent_id": 10,
            },
        ]

        mock_records.return_value = [
            _record(
                survey_record_id=1,
            ),
            _record(
                survey_record_id=2,
                status="draft",
            ),
            _record(
                survey_record_id=3,
                organization_unit_id=99,
            ),
        ]

        scope = SurveyScope(
            project_id=1,
            survey_batch_id=2,
            organization_unit_ids=(1,),
            canal_unit_ids=(10,),
        )

        records = load_scope_records(
            scope
        )

        self.assertEqual(
            tuple(
                record["survey_record_id"]
                for record in records
            ),
            (1,),
        )

    @patch(
        "services.engineering_batch_export."
        "load_scope_records"
    )
    def test_plan_groups_records_in_registry_order(
        self,
        mock_load,
    ):
        mock_load.return_value = (
            _record(
                survey_record_id=2,
                form_code="form_2_2",
            ),
            _record(
                survey_record_id=1,
                form_code="form_2_1",
            ),
        )

        with tempfile.TemporaryDirectory() as temp:
            request = BatchExportRequest(
                scope=SurveyScope(
                    project_id=1,
                    survey_batch_id=2,
                ),
                output_root=(
                    Path(temp)
                    / "成果"
                ),
            )

            plan = build_batch_export_plan(
                request
            )

        self.assertEqual(
            tuple(
                group.definition.form_code
                for group in plan.groups
            ),
            (
                "form_2_1",
                "form_2_2",
            ),
        )

        self.assertEqual(
            len(plan.records),
            2,
        )

    def test_execute_continues_after_original_failure(
        self,
    ):
        records = (
            _record(
                survey_record_id=1,
            ),
            _record(
                survey_record_id=2,
            ),
        )

        with tempfile.TemporaryDirectory() as temp:
            output_root = (
                Path(temp)
                / "2026成果"
            )

            request = BatchExportRequest(
                scope=SurveyScope(
                    project_id=1,
                    survey_batch_id=2,
                ),
                output_root=output_root,
                create_zip=True,
            )

            plan = BatchExportPlan(
                request=request,
                records=records,
                groups=(
                    FormExportGroup(
                        definition=FORM_2_1,
                        records=records,
                    ),
                ),
            )

            def fake_summary(
                definition,
                *,
                records,
                file_path,
            ):
                Path(file_path).write_text(
                    "summary",
                    encoding="utf-8",
                )

                return {
                    "file_path": file_path,
                    "exported_count": (
                        len(records)
                    ),
                }

            def fake_original(
                definition,
                *,
                survey_record_id,
                file_path,
            ):
                if survey_record_id == 2:
                    raise ValueError(
                        "模拟原表失败"
                    )

                Path(file_path).write_text(
                    "original",
                    encoding="utf-8",
                )

                return {
                    "file_path": file_path,
                }

            with patch(
                "services.engineering_batch_export."
                "export_engineering_summary",
                side_effect=fake_summary,
            ), patch(
                "services.engineering_batch_export."
                "export_engineering_original_form",
                side_effect=fake_original,
            ):
                result = execute_batch_export(
                    plan
                )

            self.assertFalse(
                result.completed
            )
            self.assertEqual(
                result.summary_success_count,
                1,
            )
            self.assertEqual(
                result.original_success_count,
                1,
            )
            self.assertTrue(
                result.manifest_path.exists()
            )
            self.assertTrue(
                (
                    output_root
                    / "导出未完成.txt"
                ).exists()
            )
            self.assertIsNone(
                result.archive_path
            )

            workbook = load_workbook(
                result.manifest_path,
                read_only=True,
            )

            try:
                record_sheet = workbook[
                    "记录清单"
                ]

                statuses = [
                    record_sheet.cell(
                        row=row,
                        column=11,
                    ).value
                    for row in range(
                        2,
                        record_sheet.max_row
                        + 1,
                    )
                ]

                self.assertEqual(
                    statuses,
                    [
                        "成功",
                        "失败",
                    ],
                )
            finally:
                workbook.close()

    def test_execute_success_can_create_zip(
        self,
    ):
        records = (
            _record(
                survey_record_id=1,
            ),
        )

        with tempfile.TemporaryDirectory() as temp:
            output_root = (
                Path(temp)
                / "2026成果"
            )

            request = BatchExportRequest(
                scope=SurveyScope(
                    project_id=1,
                    survey_batch_id=2,
                ),
                output_root=output_root,
                create_zip=True,
            )

            plan = BatchExportPlan(
                request=request,
                records=records,
                groups=(
                    FormExportGroup(
                        definition=FORM_2_1,
                        records=records,
                    ),
                ),
            )

            def fake_export(
                *args,
                **kwargs,
            ):
                file_path = Path(
                    kwargs["file_path"]
                )

                file_path.write_text(
                    "ok",
                    encoding="utf-8",
                )

                return {
                    "file_path": file_path,
                }

            with patch(
                "services.engineering_batch_export."
                "export_engineering_summary",
                side_effect=fake_export,
            ), patch(
                "services.engineering_batch_export."
                "export_engineering_original_form",
                side_effect=fake_export,
            ):
                result = execute_batch_export(
                    plan
                )

            self.assertTrue(
                result.completed
            )
            self.assertEqual(
                result.original_success_count,
                1,
            )
            self.assertIsNotNone(
                result.archive_path
            )
            self.assertTrue(
                result.archive_path.exists()
            )
            self.assertFalse(
                (
                    output_root
                    / "导出未完成.txt"
                ).exists()
            )

    def test_nonempty_output_directory_is_rejected(
        self,
    ):
        records = (
            _record(
                survey_record_id=1,
            ),
        )

        with tempfile.TemporaryDirectory() as temp:
            output_root = (
                Path(temp)
                / "2026成果"
            )

            output_root.mkdir()

            (
                output_root
                / "已有文件.txt"
            ).write_text(
                "existing",
                encoding="utf-8",
            )

            request = BatchExportRequest(
                scope=SurveyScope(
                    project_id=1,
                    survey_batch_id=2,
                ),
                output_root=output_root,
            )

            plan = BatchExportPlan(
                request=request,
                records=records,
                groups=(
                    FormExportGroup(
                        definition=FORM_2_1,
                        records=records,
                    ),
                ),
            )

            with self.assertRaisesRegex(
                ValueError,
                "不是空目录",
            ):
                execute_batch_export(
                    plan
                )


if __name__ == "__main__":
    unittest.main()
