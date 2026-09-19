import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


from forms.engineering.form_2_3 import (
    FORM_2_3,
)
from services.engineering_summary_export import (
    export_engineering_summary,
)


class GenericEngineeringSummaryExportTestCase(
    unittest.TestCase,
):
    def test_form_2_3_summary_definition_contract(
        self,
    ):
        definition = (
            FORM_2_3.summary_export_definition
        )

        self.assertIsNotNone(definition)

        assert definition is not None

        self.assertEqual(
            definition.sheet_name,
            "渡槽调查汇总",
        )

        self.assertEqual(
            len(definition.columns),
            20,
        )

        self.assertEqual(
            tuple(
                column.header
                for column
                in definition.columns[:6]
            ),
            (
                "业务编号",
                "名称",
                "基层处",
                "水管所",
                "渠系",
                "桩号",
            ),
        )

    def test_generic_export_derives_evaluation_columns_from_definition(
        self,
    ):
        summary_definition = (
            FORM_2_3.summary_export_definition
        )

        self.assertIsNotNone(
            summary_definition
        )

        assert summary_definition is not None

        summary_record = {
            "survey_record_id": 101,
            "department_name": "测试基层处",
            "office_name": "测试水管所",
            "canal_name": "测试干渠",
            "updated_at": "2026-09-16 10:00:00",
        }

        record = {
            "business_code": "1-01-01-03-001",
            "asset_name": "测试渡槽",
            "record_data": {
                "stake": "K26+100",
                "design_flow": 8.5,
                "structure_grade": "3级",
                "build_date": "2010-06",
                "renovation_date": "2021-09",
                "length": 150.0,
                "increased_flow": 10.0,
                "structure_form": "梁式渡槽",
                "section_width": 3.2,
                "section_height": 2.4,
                "trough_body_structure": "钢筋混凝土",
                "trough_wall_thickness": 0.25,
                "waterstop_form": "橡胶止水",
                "trough_bottom_elevation": -3.5,
                "span_count": 5,
                "lower_support_structure_form": "排架式",
            },
            "record_status": "completed",
            "overall_grade": "B",
            "survey_date": "2026-09-16",
            "survey_comment": "测试意见。",
        }

        inspection_results = [
            {
                "item_code": item["item_code"],
                "grade": "B",
            }
            for item
            in FORM_2_3.evaluation_items
        ]

        with tempfile.TemporaryDirectory() as directory:
            file_path = (
                Path(directory)
                / "generic_summary.xlsx"
            )

            with patch(
                "services.engineering_summary_export."
                "get_engineering_record",
                return_value=record,
            ), patch(
                "services.engineering_summary_export."
                "get_inspection_results",
                return_value=inspection_results,
            ):
                result = export_engineering_summary(
                    FORM_2_3,
                    records=[summary_record],
                    file_path=file_path,
                )

            self.assertEqual(
                result["exported_count"],
                1,
            )

            workbook = load_workbook(
                file_path
            )

            worksheet = workbook[
                "渡槽调查汇总"
            ]

            self.assertEqual(
                worksheet["A1"].value,
                "序号",
            )
            self.assertEqual(
                worksheet["C1"].value,
                "名称",
            )
            self.assertEqual(
                worksheet["G1"].value,
                "桩号",
            )
            self.assertEqual(
                worksheet["O2"].value,
                "3.2×2.4",
            )

            evaluation_start_column = (
                2
                + len(
                    summary_definition.columns
                )
            )

            for column_index in range(
                evaluation_start_column,
                (
                    evaluation_start_column
                    + len(
                        FORM_2_3.evaluation_items
                    )
                ),
            ):
                self.assertEqual(
                    worksheet.cell(
                        row=2,
                        column=column_index,
                    ).value,
                    "B",
                )

            overall_grade_column = (
                evaluation_start_column
                + len(
                    FORM_2_3.evaluation_items
                )
            )

            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=overall_grade_column,
                ).value,
                "B",
            )

            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=(
                        overall_grade_column + 7
                    ),
                ).value,
                "录入完成",
            )

            workbook.close()


if __name__ == "__main__":
    unittest.main()
