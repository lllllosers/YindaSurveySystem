import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


from forms.engineering.registry import (
    get_engineering_form_definition,
    get_engineering_form_definitions,
    get_engineering_grade_options,
)


class BatchBRegistryTemplatesTestCase(
    unittest.TestCase,
):
    EXPECTED = {
        "form_2_10": (
            "附表2.10",
            "form_2_10_V1.xlsx",
            12,
        ),
        "form_2_11": (
            "附表2.11",
            "form_2_11_V1.xlsx",
            11,
        ),
        "form_2_14": (
            "附表2.14",
            "form_2_14_V1.xlsx",
            6,
        ),
    }

    def test_batch_b_forms_are_registered_in_numeric_order(
        self,
    ):
        definitions = (
            get_engineering_form_definitions()
        )

        actual_codes = tuple(
            definition.form_code
            for definition
            in definitions
        )

        for form_code in self.EXPECTED:
            self.assertIn(
                form_code,
                actual_codes,
            )

        self.assertLess(
            actual_codes.index(
                "form_2_9"
            ),
            actual_codes.index(
                "form_2_10"
            ),
        )
        self.assertLess(
            actual_codes.index(
                "form_2_11"
            ),
            actual_codes.index(
                "form_2_12"
            ),
        )
        self.assertLess(
            actual_codes.index(
                "form_2_13"
            ),
            actual_codes.index(
                "form_2_14"
            ),
        )

    def test_each_batch_b_form_has_three_grade_contract(
        self,
    ):
        for form_code in self.EXPECTED:
            with self.subTest(
                form_code=form_code,
            ):
                self.assertEqual(
                    get_engineering_grade_options(
                        form_code
                    ),
                    (
                        "A",
                        "B",
                        "C",
                    ),
                )

    def test_templates_match_original_form_declarations(
        self,
    ):
        template_dir = (
            PROJECT_ROOT
            / "templates"
            / "excel"
        )

        for (
            form_code,
            (
                sheet_name,
                template_filename,
                evaluation_count,
            ),
        ) in self.EXPECTED.items():
            with self.subTest(
                form_code=form_code,
            ):
                definition = (
                    get_engineering_form_definition(
                        form_code
                    )
                )

                self.assertIsNotNone(
                    definition
                )
                assert definition is not None

                self.assertEqual(
                    len(
                        definition
                        .evaluation_items
                    ),
                    evaluation_count,
                )

                original = (
                    definition
                    .original_form_export_definition
                )
                summary = (
                    definition
                    .summary_export_definition
                )

                self.assertIsNotNone(
                    original
                )
                self.assertIsNotNone(
                    summary
                )
                assert original is not None

                self.assertEqual(
                    original.sheet_name,
                    sheet_name,
                )
                self.assertEqual(
                    original.template_filename,
                    template_filename,
                )

                template_path = (
                    template_dir
                    / template_filename
                )

                self.assertTrue(
                    template_path.exists()
                )

                workbook = load_workbook(
                    template_path,
                    data_only=False,
                )

                try:
                    self.assertIn(
                        sheet_name,
                        workbook.sheetnames,
                    )

                    worksheet = workbook[
                        sheet_name
                    ]

                    self.assertEqual(
                        worksheet["A1"].value,
                        sheet_name,
                    )

                    header_row = (
                        original
                        .evaluation_binding
                        .start_row
                        - 1
                    )

                    self.assertEqual(
                        worksheet[
                            f"A{header_row}"
                        ].value,
                        "项目",
                    )

                    for offset in range(
                        evaluation_count
                    ):
                        description = (
                            worksheet[
                                f"F{original.evaluation_binding.start_row + offset}"
                            ].value
                            or ""
                        )

                        self.assertIn(
                            "a ",
                            description,
                        )
                        self.assertIn(
                            "b ",
                            description,
                        )
                        self.assertIn(
                            "c ",
                            description,
                        )
                        self.assertNotIn(
                            "d ",
                            description,
                        )

                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
