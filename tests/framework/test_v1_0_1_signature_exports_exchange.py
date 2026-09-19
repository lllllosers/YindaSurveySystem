import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from openpyxl import Workbook, load_workbook

from forms.engineering.form_2_4 import FORM_2_4
from forms.engineering.registry import get_engineering_form_definitions
from services.engineering_original_form_export import export_engineering_original_form
from services.engineering_summary_export import export_engineering_summary
from services.survey_result_package import (
    RESULT_SCHEMA_VERSION,
    SUPPORTED_RESULT_SCHEMA_VERSIONS,
)


class SignatureExportContractTestCase(unittest.TestCase):
    def test_result_schema_is_2_1_and_reader_keeps_2_0_compatibility(self):
        self.assertEqual(RESULT_SCHEMA_VERSION, "2.1")
        self.assertEqual(
            SUPPORTED_RESULT_SCHEMA_VERSIONS,
            ("2.0", "2.1"),
        )

    def test_all_14_original_form_signature_cells_follow_footer_contract(self):
        definitions = get_engineering_form_definitions()
        self.assertEqual(len(definitions), 14)

        for definition in definitions:
            original = definition.original_form_export_definition
            self.assertIsNotNone(original)
            assert original is not None

            conclusion = original.conclusion_binding
            self.assertTrue(conclusion.survey_date_cell.startswith("J"))
            row = conclusion.survey_date_cell[1:]

            with self.subTest(form_code=definition.form_code):
                self.assertEqual(
                    conclusion.surveyor_signatures_cell,
                    f"B{row}",
                )
                self.assertEqual(
                    conclusion.water_office_manager_signature_cell,
                    f"D{row}",
                )
                self.assertEqual(
                    conclusion.engineering_section_chief_signature_cell,
                    f"F{row}",
                )
                self.assertEqual(
                    conclusion.department_head_signature_cell,
                    f"H{row}",
                )

    def test_generic_original_export_writes_four_signature_values(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template_dir = root / "templates" / "excel"
            template_dir.mkdir(parents=True)

            original = FORM_2_4.original_form_export_definition
            assert original is not None

            workbook = Workbook()
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.title = original.sheet_name
            workbook.save(template_dir / original.template_filename)
            workbook.close()

            record = {
                "engineering_asset_id": 10,
                "canal_id": 20,
                "business_code": "1-01-01-04-001",
                "asset_name": "签字测试倒虹吸",
                "record_data": {
                    "asset_name": "签字测试倒虹吸",
                    "stake": "CH1+000",
                },
                "survey_comment": "签字测试意见",
                "overall_grade": "B",
                "survey_date": "2026-09-19",
                "surveyor_signatures": "张三、李四",
                "water_office_manager_signature": "王五",
                "engineering_section_chief_signature": "赵六",
                "department_head_signature": "钱七",
            }

            inspections = [
                {
                    "item_code": item["item_code"],
                    "grade": "B",
                }
                for item in FORM_2_4.evaluation_items
            ]

            output_path = root / "signature_original.xlsx"

            with patch(
                "services.engineering_original_form_export.get_app_root",
                return_value=root,
            ), patch(
                "services.engineering_original_form_export.get_engineering_record",
                return_value=record,
            ), patch(
                "services.engineering_original_form_export.get_engineering_asset_detail",
                return_value={"id": 10},
            ), patch(
                "services.engineering_original_form_export.get_inspection_results",
                return_value=inspections,
            ), patch(
                "services.engineering_original_form_export.fill_original_form_ownership_header",
            ):
                export_engineering_original_form(
                    FORM_2_4,
                    survey_record_id=99,
                    file_path=output_path,
                )

            workbook = load_workbook(output_path)
            worksheet = workbook[original.sheet_name]
            row = original.conclusion_binding.survey_date_cell[1:]

            self.assertEqual(worksheet[f"B{row}"].value, "张三、李四")
            self.assertEqual(worksheet[f"D{row}"].value, "王五")
            self.assertEqual(worksheet[f"F{row}"].value, "赵六")
            self.assertEqual(worksheet[f"H{row}"].value, "钱七")
            self.assertEqual(worksheet[f"J{row}"].value, "2026-09-19")
            workbook.close()

    def test_generic_summary_exports_four_signature_columns(self):
        summary_record = {
            "survey_record_id": 101,
            "department_name": "测试基层处",
            "office_name": "测试水管所",
            "canal_name": "测试干渠",
            "updated_at": "2026-09-19 12:00:00",
        }

        record = {
            "business_code": "1-01-01-04-001",
            "asset_name": "签字测试倒虹吸",
            "record_data": {
                "asset_name": "签字测试倒虹吸",
                "stake": "CH1+000",
            },
            "record_status": "completed",
            "overall_grade": "B",
            "survey_date": "2026-09-19",
            "survey_comment": "签字测试意见",
            "surveyor_signatures": "张三、李四",
            "water_office_manager_signature": "王五",
            "engineering_section_chief_signature": "赵六",
            "department_head_signature": "钱七",
        }

        inspections = [
            {
                "item_code": item["item_code"],
                "grade": "B",
            }
            for item in FORM_2_4.evaluation_items
        ]

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "signature_summary.xlsx"

            with patch(
                "services.engineering_summary_export.get_engineering_record",
                return_value=record,
            ), patch(
                "services.engineering_summary_export.get_inspection_results",
                return_value=inspections,
            ):
                export_engineering_summary(
                    FORM_2_4,
                    records=[summary_record],
                    file_path=output_path,
                )

            workbook = load_workbook(output_path)
            worksheet = workbook[
                FORM_2_4.summary_export_definition.sheet_name
            ]

            headers = [
                worksheet.cell(row=1, column=index).value
                for index in range(1, worksheet.max_column + 1)
            ]

            for header in (
                "调查人签字",
                "水管所负责人",
                "工程科科长",
                "基层处负责人",
            ):
                self.assertIn(header, headers)

            values_by_header = {
                worksheet.cell(row=1, column=index).value:
                worksheet.cell(row=2, column=index).value
                for index in range(1, worksheet.max_column + 1)
            }

            self.assertEqual(values_by_header["调查人签字"], "张三、李四")
            self.assertEqual(values_by_header["水管所负责人"], "王五")
            self.assertEqual(values_by_header["工程科科长"], "赵六")
            self.assertEqual(values_by_header["基层处负责人"], "钱七")

            workbook.close()


if __name__ == "__main__":
    unittest.main()
