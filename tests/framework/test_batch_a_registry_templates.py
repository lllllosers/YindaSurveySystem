import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))


from forms.engineering.registry import (
    get_engineering_form_definition,
    get_engineering_form_definitions,
)


class BatchARegistryTemplatesTestCase(unittest.TestCase):
    EXPECTED = {
        "form_2_7": ("附表2.7", "form_2_7_V1.xlsx", 10),
        "form_2_8": ("附表2.8", "form_2_8_V1.xlsx", 7),
        "form_2_9": ("附表2.9", "form_2_9_V1.xlsx", 9),
        "form_2_12": ("附表2.12", "form_2_12_V1.xlsx", 10),
        "form_2_13": ("附表2.13", "form_2_13_V1.xlsx", 4),
    }

    def test_batch_a_forms_are_registered(self):
        registered = {
            definition.form_code
            for definition
            in get_engineering_form_definitions()
        }

        self.assertTrue(
            set(self.EXPECTED).issubset(registered)
        )

    def test_batch_a_templates_match_declared_contract(self):
        template_dir = PROJECT_ROOT / "templates" / "excel"

        for form_code, (sheet_name, filename, item_count) in self.EXPECTED.items():
            with self.subTest(form_code=form_code):
                definition = get_engineering_form_definition(form_code)
                self.assertIsNotNone(definition)
                assert definition is not None

                summary = definition.summary_export_definition
                original = definition.original_form_export_definition

                self.assertIsNotNone(summary)
                self.assertIsNotNone(original)
                assert summary is not None
                assert original is not None

                self.assertEqual(
                    len(definition.evaluation_items),
                    item_count,
                )
                self.assertEqual(
                    original.template_filename,
                    filename,
                )
                self.assertEqual(
                    original.sheet_name,
                    sheet_name,
                )

                template_path = template_dir / filename
                self.assertTrue(template_path.exists())

                workbook = load_workbook(template_path, data_only=False)
                try:
                    self.assertIn(sheet_name, workbook.sheetnames)
                    worksheet = workbook[sheet_name]
                    self.assertEqual(worksheet["A1"].value, sheet_name)

                    header_row = original.evaluation_binding.start_row - 1
                    self.assertEqual(
                        worksheet[f"A{header_row}"].value,
                        "项目",
                    )
                    self.assertEqual(
                        worksheet[f"E{header_row}"].value,
                        "项目类别",
                    )
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
