import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import (
    Workbook,
    load_workbook,
)


PROJECT_ROOT = (
    Path(__file__).resolve().parents[2]
)

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


from forms.engineering.form_2_4 import (
    FORM_2_4,
)
from services.engineering_original_form_export import (
    export_engineering_original_form,
)


class GenericEngineeringOriginalFormExportTestCase(
    unittest.TestCase,
):
    def test_generic_original_export_executes_declared_bindings(
        self,
    ):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)

            template_dir = (
                root
                / "templates"
                / "excel"
            )
            template_dir.mkdir(
                parents=True
            )

            template_path = (
                template_dir
                / "form_2_4_V1.xlsx"
            )

            workbook = Workbook()
            worksheet = workbook.active

            assert worksheet is not None

            worksheet.title = "附表2.4"
            workbook.save(template_path)
            workbook.close()

            record = {
                "engineering_asset_id": 10,
                "canal_id": 20,
                "business_code": (
                    "1-01-01-04-001"
                ),
                "asset_name": "测试倒虹吸",
                "record_data": {
                    "stake": "K10+500",
                    "design_flow": 8.5,
                    "structure_grade": "3级",
                    "build_date": "2010-06",
                    "renovation_date": None,
                    "length": 260.0,
                    "increased_flow": 10.0,
                    "structure_form": "埋管式",
                    "section_size": "DN2400",
                    "pipe_body_structure": (
                        "钢筋混凝土"
                    ),
                    "wall_thickness": 0.3,
                    "waterstop_form": (
                        "橡胶止水"
                    ),
                    "channel_bottom_elevation": (
                        1680.25
                    ),
                },
                "survey_comment": "测试意见。",
                "overall_grade": "C",
                "survey_date": "2026-09-16",
            }

            inspections = [
                {
                    "item_code": (
                        item["item_code"]
                    ),
                    "grade": "B",
                }
                for item
                in FORM_2_4.evaluation_items
            ]

            output_path = (
                root
                / "export.xlsx"
            )

            with patch(
                "services."
                "engineering_original_form_export."
                "get_app_root",
                return_value=root,
            ), patch(
                "services."
                "engineering_original_form_export."
                "get_engineering_record",
                return_value=record,
            ), patch(
                "services."
                "engineering_original_form_export."
                "get_engineering_asset_detail",
                return_value={"id": 10},
            ), patch(
                "services."
                "engineering_original_form_export."
                "get_inspection_results",
                return_value=inspections,
            ), patch(
                "services."
                "engineering_original_form_export."
                "fill_original_form_ownership_header",
            ) as fill_header:
                result = (
                    export_engineering_original_form(
                        FORM_2_4,
                        survey_record_id=99,
                        file_path=output_path,
                    )
                )

            self.assertEqual(
                result["survey_record_id"],
                99,
            )
            self.assertEqual(
                result["business_code"],
                "1-01-01-04-001",
            )
            self.assertEqual(
                result["asset_name"],
                "测试倒虹吸",
            )

            fill_header.assert_called_once()

            workbook = load_workbook(
                output_path,
                data_only=False,
            )

            worksheet = workbook[
                "附表2.4"
            ]

            self.assertEqual(
                worksheet["B5"].value,
                "测试倒虹吸",
            )
            self.assertEqual(
                worksheet["D7"].value,
                "DN2400",
            )
            self.assertEqual(
                worksheet["B8"].value,
                1680.25,
            )

            for row_number in range(
                10,
                22,
            ):
                self.assertEqual(
                    worksheet[
                        f"E{row_number}"
                    ].value,
                    "B",
                )

            self.assertEqual(
                worksheet["C22"].value,
                "测试意见。",
            )
            self.assertEqual(
                worksheet["J22"].value,
                "C",
            )
            self.assertEqual(
                worksheet["J23"].value,
                "2026-09-16",
            )

            self.assertEqual(
                worksheet.print_area,
                "'附表2.4'!$A$1:$J$25",
            )
            self.assertEqual(
                worksheet.page_setup.orientation,
                "landscape",
            )
            self.assertEqual(
                worksheet.page_setup.fitToWidth,
                1,
            )
            self.assertEqual(
                worksheet.page_setup.fitToHeight,
                1,
            )
            self.assertFalse(
                worksheet.sheet_view
                .showGridLines
            )

            workbook.close()


if __name__ == "__main__":
    unittest.main()
