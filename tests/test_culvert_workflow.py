import sys
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from workflow_test_support import (
    EngineeringWorkflowTestCaseBase,
)

from services.culvert_evaluation import (
    CULVERT_EVALUATION_ITEMS,
)

from openpyxl import (
    load_workbook,
)

from services.culvert_export import (
    export_culvert_original_form,
    export_culvert_summary,
)


class CulvertWorkflowTestCase(
    EngineeringWorkflowTestCaseBase,
):
    """
    附表2.6涵洞（暗涵）
    核心业务回归测试。

    当前阶段验证：
    - 公共 point engineering 创建；
    - 单桩号工程身份；
    - 重复保护；
    - 公共读取和修改；
    - 公共删除；
    - 附表2.6完成调查规则。
    """

    FORM_CODE = "form_2_6"

    TEST_DB_FILENAME = "test_culvert.db"

    PROJECT_NAME = "涵洞自动测试项目"

    PROJECT_SHORT_NAME = "涵洞测试"

    BATCH_NAME = "涵洞自动测试批次"

    BATCH_CODE = "CULVERT_TEST_2026"

    # =========================================================
    # 测试数据辅助
    # =========================================================

    def make_complete_record_data(
        self,
        *,
        asset_name="测试涵洞",
        stake="K30+500",
        stake_value=30500.0,
    ):
        return {
            "asset_name": asset_name,
            "stake": stake,
            "stake_value": stake_value,
            "design_flow": 6.5,
            "structure_grade": "3级",
            "build_date": "2012-06",
            "renovation_date": None,
            "length": 85.0,
            "increased_flow": 7.8,
            "structure_form": ("钢筋混凝土箱涵"),
            "main_structure_material": ("钢筋混凝土"),
            "concrete_strength": "C30",
            "cover_thickness": 40.0,
            "soil_cover_thickness": 2.5,
            "channel_width": 3.2,
            "channel_depth": 2.8,
            "channel_bottom_elevation": (1685.35),
        }

    def make_complete_inspection_results(
        self,
        grade="B",
    ):
        return [
            {
                "item_code": (item["item_code"]),
                "category": (item["category"]),
                "item_name": (item["item_name"]),
                "grade": grade,
                "description": None,
                "remark": None,
            }
            for item in (CULVERT_EVALUATION_ITEMS)
        ]

    def create_culvert_draft(
        self,
        *,
        business_code=("1-01-01-06-001"),
        asset_name="测试涵洞",
        stake="K30+500",
        stake_value=30500.0,
        record_data=None,
        inspection_results=None,
        survey_date=None,
        overall_grade=None,
        survey_comment=None,
    ):
        if record_data is None:
            record_data = {
                "asset_name": asset_name,
                "stake": stake,
                "stake_value": stake_value,
                "design_flow": 6.5,
                "length": 85.0,
            }

        return database.create_engineering_survey(
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(self.form_version["id"]),
            asset_name=asset_name,
            asset_type=(self.form_version["asset_type"]),
            organization_unit_id=(self.office_id),
            canal_unit_id=self.canal_id,
            business_code=business_code,
            record_data=record_data,
            single_stake_text=stake,
            single_stake_value=(stake_value),
            inspection_results=(inspection_results),
            survey_date=survey_date,
            overall_grade=overall_grade,
            survey_comment=survey_comment,
        )

    # =========================================================
    # 1. point engineering 创建
    # =========================================================

    def test_create_culvert_point_draft(
        self,
    ):
        result = self.create_culvert_draft()

        survey_record_id = int(result["survey_record_id"])

        record = database.get_point_engineering_record(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
        )

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            record["asset_type"],
            "culvert",
        )

        self.assertEqual(
            record["asset_name"],
            "测试涵洞",
        )

        self.assertEqual(
            record["single_stake_text"],
            "K30+500",
        )

        self.assertAlmostEqual(
            float(record["single_stake_value"]),
            30500.0,
        )

        self.assertEqual(
            record["record_status"],
            "draft",
        )

        records = database.get_engineering_survey_query_records(
            project_id=self.project_id,
            survey_batch_id=(self.batch_id),
            form_code="form_2_6",
        )

        self.assertEqual(
            len(records),
            1,
        )

        self.assertEqual(
            records[0]["engineering_position"],
            "K30+500",
        )

    # =========================================================
    # 2. 同批次同桩号重复保护
    # =========================================================

    def test_duplicate_culvert_is_rejected(
        self,
    ):
        self.create_culvert_draft(
            business_code=("1-01-01-06-001"),
            asset_name="测试涵洞一",
            stake="K31+000",
            stake_value=31000.0,
        )

        with self.assertRaises(ValueError) as context:
            self.create_culvert_draft(
                business_code=("1-01-01-06-002"),
                asset_name=("另一个涵洞名称"),
                stake="K31+000",
                stake_value=31000.0,
            )

        self.assertIn(
            "工程调查记录",
            str(context.exception),
        )

    # =========================================================
    # 3. 公共读取和修改
    # =========================================================

    def test_get_and_update_culvert_record(
        self,
    ):
        result = self.create_culvert_draft()

        survey_record_id = int(result["survey_record_id"])

        engineering_asset_id = int(result["engineering_asset_id"])

        record = database.get_point_engineering_record(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
        )

        self.assertIsNotNone(record)

        assert record is not None

        updated_data = dict(record["record_data"])

        updated_data.update(
            {
                "asset_name": ("修改后的涵洞"),
                "stake": "K30+800",
                "stake_value": 30800.0,
                "design_flow": 7.2,
                "length": 90.0,
            }
        )

        database.update_point_engineering_survey(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
            asset_name="修改后的涵洞",
            record_data=updated_data,
            single_stake_text="K30+800",
            single_stake_value=30800.0,
        )

        updated_record = database.get_point_engineering_record(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
        )

        self.assertIsNotNone(updated_record)

        assert updated_record is not None

        self.assertEqual(
            updated_record["asset_name"],
            "修改后的涵洞",
        )

        self.assertEqual(
            updated_record["single_stake_text"],
            "K30+800",
        )

        self.assertEqual(
            updated_record["record_data"]["design_flow"],
            7.2,
        )

        self.assertEqual(
            int(updated_record["engineering_asset_id"]),
            engineering_asset_id,
        )

    # =========================================================
    # 4. 公共删除
    # =========================================================

    def test_delete_culvert_record(
        self,
    ):
        result = self.create_culvert_draft()

        survey_record_id = int(result["survey_record_id"])

        engineering_asset_id = int(result["engineering_asset_id"])

        delete_result = database.delete_engineering_survey_record(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
        )

        self.assertTrue(delete_result["asset_deleted"])

        self.assertEqual(
            delete_result["engineering_asset_id"],
            engineering_asset_id,
        )

        with database.get_connection() as connection:
            record_count = connection.execute(
                """
                    SELECT COUNT(*) AS count
                    FROM survey_records
                    WHERE id = ?
                    """,
                (survey_record_id,),
            ).fetchone()["count"]

            asset_count = connection.execute(
                """
                    SELECT COUNT(*) AS count
                    FROM engineering_assets
                    WHERE id = ?
                    """,
                (engineering_asset_id,),
            ).fetchone()["count"]

        self.assertEqual(
            record_count,
            0,
        )

        self.assertEqual(
            asset_count,
            0,
        )

    # =========================================================
    # 5. 正常完成调查
    # =========================================================

    def test_complete_culvert_record(
        self,
    ):
        result = self.create_culvert_draft(
            business_code=("1-01-01-06-010"),
            record_data=(self.make_complete_record_data()),
            inspection_results=(self.make_complete_inspection_results()),
            survey_date="2026-09-15",
            overall_grade="B",
            survey_comment=("自动测试调查意见。"),
        )

        survey_record_id = int(result["survey_record_id"])

        complete_result = database.complete_culvert_record(survey_record_id)

        self.assertEqual(
            complete_result["inspection_count"],
            11,
        )

        record = database.get_point_engineering_record(
            survey_record_id=(survey_record_id),
            form_code="form_2_6",
        )

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            record["record_status"],
            "completed",
        )

    # =========================================================
    # 6. 11项评价必须全部完成
    # =========================================================

    def test_completion_requires_all_11_items(
        self,
    ):
        inspections = (self.make_complete_inspection_results())[:-1]

        result = self.create_culvert_draft(
            business_code=("1-01-01-06-011"),
            asset_name=("评价不完整涵洞"),
            stake="K31+500",
            stake_value=31500.0,
            record_data=(
                self.make_complete_record_data(
                    asset_name=("评价不完整涵洞"),
                    stake="K31+500",
                    stake_value=31500.0,
                )
            ),
            inspection_results=inspections,
            survey_date="2026-09-15",
            overall_grade="B",
            survey_comment="自动测试。",
        )

        with self.assertRaisesRegex(
            ValueError,
            "全部11项",
        ):
            database.complete_culvert_record(int(result["survey_record_id"]))

    # =========================================================
    # 7. 正式基本信息必须完整
    # =========================================================

    def test_completion_requires_basic_fields(
        self,
    ):
        record_data = self.make_complete_record_data(
            asset_name=("缺少覆土厚度涵洞"),
            stake="K32+000",
            stake_value=32000.0,
        )

        record_data["soil_cover_thickness"] = None

        result = self.create_culvert_draft(
            business_code=("1-01-01-06-012"),
            asset_name=("缺少覆土厚度涵洞"),
            stake="K32+000",
            stake_value=32000.0,
            record_data=record_data,
            inspection_results=(self.make_complete_inspection_results()),
            survey_date="2026-09-15",
            overall_grade="B",
            survey_comment="自动测试。",
        )

        with self.assertRaisesRegex(
            ValueError,
            "覆土厚度",
        ):
            database.complete_culvert_record(int(result["survey_record_id"]))

    def test_export_culvert_original_form(
        self,
    ):
        """
        附表2.6正式原表
        应按模板坐标准确写入。
        """

        result = self.create_culvert_draft(
            business_code=("1-01-01-06-020"),
            asset_name=("正式导出测试涵洞"),
            stake="K33+500",
            stake_value=33500.0,
            record_data=(
                self.make_complete_record_data(
                    asset_name=("正式导出测试涵洞"),
                    stake="K33+500",
                    stake_value=33500.0,
                )
            ),
            inspection_results=(self.make_complete_inspection_results(grade="B")),
            survey_date=("2026-09-15"),
            overall_grade="C",
            survey_comment=("涵洞正式导出测试。"),
        )

        survey_record_id = int(result["survey_record_id"])

        output_path = Path(self.temp_directory.name) / "form_2_6_export.xlsx"

        export_culvert_original_form(
            survey_record_id=(survey_record_id),
            file_path=output_path,
        )

        self.assertTrue(output_path.exists())

        workbook = load_workbook(
            output_path,
            data_only=False,
        )

        worksheet = workbook["附表2.6"]

        # =========================
        # 基本信息
        # =========================

        self.assertEqual(
            worksheet["B5"].value,
            "正式导出测试涵洞",
        )

        self.assertEqual(
            worksheet["H5"].value,
            "K33+500",
        )

        self.assertEqual(
            worksheet["J5"].value,
            6.5,
        )

        self.assertEqual(
            worksheet["D7"].value,
            "钢筋混凝土",
        )

        self.assertEqual(
            worksheet["F7"].value,
            "C30",
        )

        self.assertEqual(
            worksheet["H7"].value,
            40.0,
        )

        self.assertEqual(
            worksheet["J7"].value,
            2.5,
        )

        self.assertEqual(
            worksheet["B8"].value,
            3.2,
        )

        self.assertEqual(
            worksheet["D8"].value,
            2.8,
        )

        self.assertEqual(
            worksheet["F8"].value,
            1685.35,
        )

        # =========================
        # 11项评价
        # =========================

        for row_number in range(
            10,
            21,
        ):
            self.assertEqual(
                worksheet[f"E{row_number}"].value,
                "B",
            )

        # =========================
        # 调查结论
        # =========================

        self.assertEqual(
            worksheet["C21"].value,
            "涵洞正式导出测试。",
        )

        self.assertEqual(
            worksheet["J21"].value,
            "C",
        )

        self.assertEqual(
            worksheet["J22"].value,
            "2026-09-15",
        )

        # 签字栏保持空白。
        self.assertIsNone(worksheet["B22"].value)

        # 正式原表注释。
        self.assertEqual(
            worksheet["A23"].value,
            ("注：涵洞（暗涵）其它结构指" "进出口渐变段、洞脸、洞顶、止水等。"),
        )

        workbook.close()

    def test_export_culvert_summary(
        self,
    ):
        """
        附表2.6详细汇总应包含
        正式基本信息、11项评价和调查结论。
        """

        result = (
            self.create_culvert_draft(
                business_code=(
                    "1-01-01-06-030"
                ),
                asset_name=(
                    "详细汇总测试涵洞"
                ),
                stake="K35+600",
                stake_value=35600.0,
                record_data=(
                    self.make_complete_record_data(
                        asset_name=(
                            "详细汇总测试涵洞"
                        ),
                        stake="K35+600",
                        stake_value=35600.0,
                    )
                ),
                inspection_results=(
                    self.make_complete_inspection_results(
                        grade="C"
                    )
                ),
                survey_date=(
                    "2026-09-15"
                ),
                overall_grade="B",
                survey_comment=(
                    "涵洞详细汇总测试。"
                ),
            )
        )

        survey_record_id = int(
            result[
                "survey_record_id"
            ]
        )

        records = (
            database
            .get_engineering_survey_query_records(
                project_id=(
                    self.project_id
                ),
                survey_batch_id=(
                    self.batch_id
                ),
                form_code=(
                    "form_2_6"
                ),
            )
        )

        target_records = [
            record
            for record in records
            if (
                record[
                    "survey_record_id"
                ]
                == survey_record_id
            )
        ]

        self.assertEqual(
            len(target_records),
            1,
        )

        output_path = (
            Path(
                self.temp_directory.name
            )
            / "culvert_summary.xlsx"
        )

        export_result = (
            export_culvert_summary(
                records=target_records,
                file_path=output_path,
            )
        )

        self.assertEqual(
            export_result[
                "exported_count"
            ],
            1,
        )

        self.assertTrue(
            output_path.exists()
        )

        workbook = load_workbook(
            output_path
        )

        worksheet = workbook[
            "涵洞（暗涵）调查汇总"
        ]

        # =========================
        # 表头
        # =========================

        self.assertEqual(
            worksheet["A1"].value,
            "序号",
        )

        self.assertEqual(
            worksheet["G1"].value,
            "桩号",
        )

        self.assertEqual(
            worksheet["O1"].value,
            "主构建筑材料",
        )

        self.assertEqual(
            worksheet["Q1"].value,
            "钢筋保护层厚度",
        )

        self.assertEqual(
            worksheet["U1"].value,
            "渠底高程（m）",
        )

        # =========================
        # 基本信息
        # =========================

        self.assertEqual(
            worksheet["B2"].value,
            "1-01-01-06-030",
        )

        self.assertEqual(
            worksheet["C2"].value,
            "详细汇总测试涵洞",
        )

        self.assertEqual(
            worksheet["G2"].value,
            "K35+600",
        )

        self.assertEqual(
            worksheet["H2"].value,
            6.5,
        )

        self.assertEqual(
            worksheet["O2"].value,
            "钢筋混凝土",
        )

        self.assertEqual(
            worksheet["P2"].value,
            "C30",
        )

        self.assertEqual(
            worksheet["Q2"].value,
            40.0,
        )

        self.assertEqual(
            worksheet["R2"].value,
            2.5,
        )

        self.assertEqual(
            worksheet["S2"].value,
            3.2,
        )

        self.assertEqual(
            worksheet["T2"].value,
            2.8,
        )

        self.assertEqual(
            worksheet["U2"].value,
            1685.35,
        )

        # =========================
        # 11项评价
        # V ～ AF
        # =========================

        for column in range(
            22,
            33,
        ):
            self.assertEqual(
                worksheet.cell(
                    row=2,
                    column=column,
                ).value,
                "C",
            )

        # =========================
        # 调查结论
        # AG ～ AK
        # =========================

        self.assertEqual(
            worksheet["AG2"].value,
            "B",
        )

        self.assertEqual(
            worksheet["AH2"].value,
            "2026-09-15",
        )

        self.assertEqual(
            worksheet["AI2"].value,
            "涵洞详细汇总测试。",
        )

        self.assertEqual(
            worksheet["AJ2"].value,
            "草稿",
        )

        workbook.close()

if __name__ == "__main__":
    unittest.main()
