import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(
        0,
        str(SRC_DIR),
    )


import database

from forms.engineering.form_2_5 import (
    FORM_2_5,
)

from forms.engineering.persistence import (
    complete_engineering_record,
    create_engineering_record,
    get_engineering_record,
    update_engineering_record,
)

from services.tunnel_export import (
    export_tunnel_original_form,
    export_tunnel_summary,
)

from tests.workflow_test_support import (
    EngineeringWorkflowTestCaseBase,
)


class TunnelWorkflowTestCase(
    EngineeringWorkflowTestCaseBase,
):
    """
    附表2.5隧洞核心业务回归测试。

    本文件验证正式迁移后的真实业务链：

    FORM_2_5
        -> generic persistence
        -> range database API
        -> completion
        -> export
    """

    FORM_CODE = FORM_2_5.form_code

    TEST_DB_FILENAME = "test_tunnel.db"

    PROJECT_NAME = "附表2.5自动测试项目"

    PROJECT_SHORT_NAME = "2.5测试"

    BATCH_NAME = "附表2.5自动测试批次"

    BATCH_CODE = "FORM_2_5_TEST"

    BATCH_START_DATE = "2026-09-01"

    BATCH_END_DATE = "2026-12-31"

    DEPARTMENT_NAME = "测试基层处"
    DEPARTMENT_CODE = "01"

    OFFICE_NAME = "测试水管所"
    OFFICE_CODE = "01"

    CANAL_NAME = "测试总干渠"
    CANAL_LEVEL = "01"
    CANAL_DESCRIPTION = "自动测试"

    BUSINESS_CODE = "1-01-01-" f"{FORM_2_5.business_type_code}" "-001"

    def make_complete_record_data(
        self,
        *,
        asset_name="测试隧洞",
        start_stake="K20+000",
        start_stake_value=20000.0,
        end_stake="K21+200",
        end_stake_value=21200.0,
    ):
        return {
            "asset_name": asset_name,
            "start_stake": start_stake,
            "start_stake_value": (start_stake_value),
            "end_stake": end_stake,
            "end_stake_value": (end_stake_value),
            "design_flow": 8.5,
            "structure_grade": "3级",
            "build_date": "2010-06",
            "renovation_date": None,
            "length": 1200.0,
            "increased_flow": 10.0,
            "lining_form": ("钢筋混凝土衬砌"),
            "lining_thickness": 0.30,
            "concrete_strength": "C30",
            "inlet_outlet_bottom_elevation": ("1680.25"),
            "longitudinal_slope": 1.5,
            "section_form": "城门洞型",
            "section_width": 3.0,
            "section_height": 3.5,
            "cover_thickness": 0.05,
        }

    def make_complete_inspection_results(
        self,
        grade="A",
    ):
        return [
            {
                "item_code": (item["item_code"]),
                "category": (item["category"]),
                "item_name": (item["item_name"]),
                "grade": grade,
                "description": (item["standards"][grade]),
                "remark": None,
            }
            for item in FORM_2_5.evaluation_items
        ]

    def make_complete_payload(
        self,
        *,
        asset_name="测试隧洞",
        start_stake="K20+000",
        start_stake_value=20000.0,
        end_stake="K21+200",
        end_stake_value=21200.0,
        design_flow=8.5,
    ):
        record_data = self.make_complete_record_data(
            asset_name=asset_name,
            start_stake=start_stake,
            start_stake_value=(start_stake_value),
            end_stake=end_stake,
            end_stake_value=(end_stake_value),
        )

        record_data["design_flow"] = design_flow

        return {
            "asset_name": asset_name,
            "record_data": (record_data),
            "position": {
                "kind": "range",
                "start_stake_text": (start_stake),
                "start_stake_value": (start_stake_value),
                "end_stake_text": (end_stake),
                "end_stake_value": (end_stake_value),
            },
            "inspection_results": (self.make_complete_inspection_results()),
            "survey_date": ("2026-09-14"),
            "overall_grade": "A",
            "survey_comment": ("隧洞工程状况良好。"),
        }

    def create_tunnel_draft(
        self,
        *,
        business_code=None,
        payload=None,
    ):
        if business_code is None:
            business_code = self.BUSINESS_CODE

        if payload is None:
            payload = self.make_complete_payload()

        return create_engineering_record(
            FORM_2_5,
            project_id=self.project_id,
            survey_batch_id=self.batch_id,
            form_version_id=(self.form_version["id"]),
            organization_unit_id=(self.office_id),
            canal_unit_id=(self.canal_id),
            business_code=business_code,
            payload=payload,
        )

    def test_create_tunnel_range_draft(
        self,
    ):
        payload = self.make_complete_payload()

        result = self.create_tunnel_draft(payload=payload)

        survey_record_id = int(result["survey_record_id"])

        engineering_asset_id = int(result["engineering_asset_id"])

        asset = database.get_engineering_asset_detail(engineering_asset_id)

        self.assertIsNotNone(asset)

        assert asset is not None

        self.assertEqual(
            asset["asset_type"],
            FORM_2_5.asset_type,
        )

        self.assertIsNone(asset["single_stake_text"])

        self.assertEqual(
            asset["start_stake_text"],
            "K20+000",
        )

        self.assertAlmostEqual(
            float(asset["start_stake_value"]),
            20000.0,
        )

        self.assertEqual(
            asset["end_stake_text"],
            "K21+200",
        )

        self.assertAlmostEqual(
            float(asset["end_stake_value"]),
            21200.0,
        )

        record = get_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
        )

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            record["record_status"],
            "draft",
        )

        self.assertEqual(
            record["record_data"]["design_flow"],
            8.5,
        )

        self.assertEqual(
            record["record_data"]["length"],
            1200.0,
        )

    def test_duplicate_tunnel_range_is_rejected(
        self,
    ):
        self.create_tunnel_draft()

        duplicate_payload = self.make_complete_payload(asset_name=("另一个隧洞"))

        with self.assertRaises(ValueError) as context:
            self.create_tunnel_draft(
                business_code=("1-01-01-05-002"),
                payload=(duplicate_payload),
            )

        self.assertIn(
            "工程调查记录",
            str(context.exception),
        )

    def test_get_and_update_tunnel_record(
        self,
    ):
        result = self.create_tunnel_draft()

        survey_record_id = int(result["survey_record_id"])

        updated_payload = self.make_complete_payload(
            asset_name=("修改后的隧洞"),
            start_stake=("K20+500"),
            start_stake_value=(20500.0),
            end_stake=("K21+800"),
            end_stake_value=(21800.0),
            design_flow=9.2,
        )

        updated_payload["record_data"]["length"] = 1300.0

        update_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
            payload=(updated_payload),
        )

        updated_record = get_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
        )

        self.assertIsNotNone(updated_record)

        assert updated_record is not None

        self.assertEqual(
            updated_record["asset_name"],
            "修改后的隧洞",
        )

        self.assertEqual(
            updated_record["start_stake_text"],
            "K20+500",
        )

        self.assertEqual(
            updated_record["end_stake_text"],
            "K21+800",
        )

        self.assertEqual(
            updated_record["record_data"]["design_flow"],
            9.2,
        )

        self.assertEqual(
            updated_record["record_data"]["length"],
            1300.0,
        )

        self.assertEqual(
            updated_record["business_code"],
            self.BUSINESS_CODE,
        )

        self.assertEqual(
            updated_record["office_id"],
            self.office_id,
        )

        self.assertEqual(
            updated_record["canal_id"],
            self.canal_id,
        )

    def test_completion_requires_all_evaluations(
        self,
    ):
        payload = self.make_complete_payload()

        payload["inspection_results"] = payload["inspection_results"][:-1]

        result = self.create_tunnel_draft(payload=payload)

        survey_record_id = int(result["survey_record_id"])

        with self.assertRaises(ValueError) as context:
            complete_engineering_record(
                FORM_2_5,
                survey_record_id=(survey_record_id),
                payload=payload,
            )

        self.assertIn(
            "未完成评价",
            str(context.exception),
        )

        record = get_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
        )

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            record["record_status"],
            "draft",
        )

    def test_complete_tunnel_through_framework(
        self,
    ):
        payload = self.make_complete_payload()

        result = self.create_tunnel_draft(payload=payload)

        survey_record_id = int(result["survey_record_id"])

        complete_result = complete_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
            payload=payload,
        )

        self.assertEqual(
            complete_result["inspection_count"],
            len(FORM_2_5.evaluation_items),
        )

        record = get_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
        )

        self.assertIsNotNone(record)

        assert record is not None

        self.assertEqual(
            record["record_status"],
            "completed",
        )

    def test_update_completed_tunnel_keeps_completed_status(
        self,
    ):
        payload = self.make_complete_payload()

        result = self.create_tunnel_draft(payload=payload)

        survey_record_id = int(result["survey_record_id"])

        complete_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
            payload=payload,
        )

        updated_payload = self.make_complete_payload(design_flow=9.5)

        update_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
            payload=(updated_payload),
        )

        updated_record = get_engineering_record(
            FORM_2_5,
            survey_record_id=(survey_record_id),
        )

        self.assertIsNotNone(updated_record)

        assert updated_record is not None

        self.assertEqual(
            updated_record["record_status"],
            "completed",
        )

        self.assertEqual(
            updated_record["record_data"]["design_flow"],
            9.5,
        )

    def test_export_tunnel_summary(
        self,
    ):
        payload = self.make_complete_payload()

        self.create_tunnel_draft(payload=payload)

        records = database.get_engineering_survey_query_records(
            project_id=(self.project_id),
            survey_batch_id=(self.batch_id),
            form_code=(FORM_2_5.form_code),
        )

        output_path = self.temp_data_dir / "tunnel_summary.xlsx"

        result = export_tunnel_summary(
            records=records,
            file_path=(output_path),
        )

        self.assertEqual(
            result["exported_count"],
            1,
        )

        self.assertTrue(output_path.exists())

        workbook = load_workbook(
            output_path,
            data_only=True,
        )

        worksheet = workbook["隧洞调查汇总"]

        self.assertEqual(
            worksheet["B2"].value,
            self.BUSINESS_CODE,
        )

        self.assertEqual(
            worksheet["C2"].value,
            "测试隧洞",
        )

        self.assertEqual(
            worksheet["G2"].value,
            "K20+000",
        )

        self.assertEqual(
            worksheet["H2"].value,
            "K21+200",
        )

        self.assertEqual(
            worksheet["U2"].value,
            "3.0*3.5",
        )

        self.assertEqual(
            worksheet["W2"].value,
            "A",
        )

        self.assertEqual(
            worksheet["AJ2"].value,
            "A",
        )

        workbook.close()

    def test_export_tunnel_original_form(
        self,
    ):
        payload = self.make_complete_payload()

        result = self.create_tunnel_draft(payload=payload)

        survey_record_id = int(result["survey_record_id"])

        output_path = self.temp_data_dir / "tunnel_original.xlsx"

        export_tunnel_original_form(
            survey_record_id=(survey_record_id),
            file_path=(output_path),
        )

        self.assertTrue(output_path.exists())

        workbook = load_workbook(
            output_path,
            data_only=True,
        )

        worksheet = workbook["附表2.5"]

        self.assertEqual(
            worksheet["B5"].value,
            "测试隧洞",
        )

        self.assertEqual(
            worksheet["F5"].value,
            "K20+000～K21+200",
        )

        self.assertEqual(
            worksheet["J5"].value,
            8.5,
        )

        self.assertEqual(
            worksheet["B7"].value,
            "钢筋混凝土衬砌",
        )

        self.assertEqual(
            worksheet["E8"].value,
            "3.0*3.5",
        )

        self.assertEqual(
            worksheet["E10"].value,
            "A",
        )

        self.assertEqual(
            worksheet["J23"].value,
            "A",
        )

        self.assertEqual(
            worksheet["J24"].value,
            "2026-09-14",
        )

        workbook.close()


if __name__ == "__main__":
    unittest.main()
