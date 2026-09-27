from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.db.session import SessionLocal
from app.main import app
from app.models.auth import User
from app.services import original_form_export
from tests.auth_helpers import cleanup_test_user, create_test_user, login_client
from forms.engineering.registry import get_engineering_form_definitions, get_engineering_form_definition


def _snapshot():
    return NS(
        departments=[NS(stable_uid="dept", name="一处")],
        offices=[NS(stable_uid="office", name="一所", parent_name="一处")],
        canals=[
            NS(stable_uid="main", master_key="main", parent_master_key=None, canal_level="01", name="主干渠"),
            NS(stable_uid="branch", master_key="branch", parent_master_key="main", canal_level="03", name="第一支渠"),
        ],
    )


def _data(form_code: str):
    definition = get_engineering_form_definition(form_code)
    record_data = {field.key: "测试值" for field in definition.fields}
    record_data.update(start_stake="1+000", end_stake="1+100")
    record = NS(
        survey_record_uid="a" * 32,
        form_code=form_code,
        organization_unit_uid="office",
        canal_unit_uid="branch",
        payload_json={
            "record_data": record_data,
            "overall_grade": "B",
            "survey_comment": "调查意见",
            "surveyor_signatures": "调查人员",
            "survey_date": "2026-09-27",
        },
    )
    asset = NS(
        asset_name="示例工程",
        business_code="业务-001",
        payload_json={"start_stake_text": "1+000", "end_stake_text": "1+100"},
    )
    inspections = [NS(item_code=definition.evaluation_items[0]["item_code"], payload_json={"grade": "C"})]
    return definition, record, asset, inspections


def test_all_fourteen_form_definitions_have_templates():
    definitions = get_engineering_form_definitions()
    assert len(definitions) == 14
    for definition in definitions:
        export = definition.original_form_export_definition
        assert export is not None
        assert (original_form_export.Path(original_form_export.REPOSITORY_ROOT) / "templates" / "excel" / export.template_filename).is_file()


@pytest.mark.parametrize("form_code", ["form_2_1", "form_2_4", "form_2_10"])
def test_central_record_renders_real_workbook(form_code):
    definition, record, asset, inspections = _data(form_code)
    with patch.object(original_form_export, "get_snapshot", return_value=_snapshot()):
        content, filename = original_form_export.export_central_original_form(record, asset, inspections)
    assert definition.form_name in filename
    workbook = load_workbook(BytesIO(content))
    export = definition.original_form_export_definition
    sheet = workbook[export.sheet_name]
    assert sheet["A3"].value == "一"
    assert sheet["C3"].value == "一"
    assert sheet["E3"].value == "主"
    assert sheet["G3"].value == "第一"
    assert sheet["J3"].value == "业务-001"
    assert sheet[export.conclusion_binding.overall_grade_cell].value == "B"
    assert sheet[export.conclusion_binding.survey_comment_cell].value == "调查意见"
    assert sheet[export.conclusion_binding.surveyor_signatures_cell].value == "调查人员"
    evaluation = export.evaluation_binding
    assert sheet[f"{evaluation.column}{evaluation.start_row}"].value == "C"
    assert sheet.print_area.replace("$", "").endswith(export.print_settings.print_area)
    assert sheet.page_setup.orientation == export.print_settings.orientation
    assert sheet.page_setup.fitToWidth == export.print_settings.fit_to_width
    first = next(item for item in export.field_bindings if item.binding.source == "record_data" and item.binding.formatter is None)
    assert sheet[first.cell].value == "测试值"
    if form_code == "form_2_1":
        start = next(item for item in export.field_bindings if item.binding.source == "record" and "start_stake_text" in item.binding.keys)
        end = next(item for item in export.field_bindings if item.binding.source == "record" and "end_stake_text" in item.binding.keys)
        assert "1+000" in sheet[start.cell].value
        assert "1+100" in sheet[end.cell].value
    workbook.close()


def test_api_scoped_export_does_not_render_foreign_record():
    user = create_test_user("viewer")
    with SessionLocal() as db:
        row = db.scalar(select(User).where(User.user_uid == user.user_uid))
        row.office_scope_uid = "b" * 32
        db.commit()
    client = TestClient(app)
    login_client(client, user.username)
    _, record, asset, inspections = _data("form_2_4")
    try:
        with patch("app.api.routes.central_records.central_record_service.get_record", return_value=(record, asset, inspections, [])):
            response = client.get(f"/api/v1/central-records/{record.survey_record_uid}/original-form.xlsx")
        assert response.status_code == 404
    finally:
        cleanup_test_user(user.user_uid)


def test_api_returns_reopenable_xlsx():
    user = create_test_user("viewer")
    client = TestClient(app)
    login_client(client, user.username)
    _, record, asset, inspections = _data("form_2_4")
    try:
        with patch("app.api.routes.central_records.central_record_service.get_record", return_value=(record, asset, inspections, [])), patch.object(original_form_export, "get_snapshot", return_value=_snapshot()):
            response = client.get(f"/api/v1/central-records/{record.survey_record_uid}/original-form.xlsx")
        assert response.status_code == 200, response.text
        assert "filename*=UTF-8''" in response.headers["content-disposition"]
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active["J3"].value == "业务-001"
        workbook.close()
    finally:
        cleanup_test_user(user.user_uid)
