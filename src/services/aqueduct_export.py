from openpyxl import load_workbook

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
)

from forms.engineering.form_2_3 import (
    FORM_2_3,
)
from forms.engineering.formatters import (
    format_dimension_pair,
)
from forms.engineering.persistence import (
    get_engineering_record,
)

from services.engineering_summary_export import (
    export_engineering_summary,
)
from services.original_form_export_common import (
    fill_original_form_ownership_header,
)


def _format_section_size(
    record_data,
):
    """
    将内部拆分保存的断面宽、高
    组合为正式字段“断面尺寸（宽×高）”。

    该薄包装仅供当前正式原表导出调用；
    实际格式化规则与详细汇总共用同一纯函数。
    """

    return format_dimension_pair(
        record_data.get("section_width"),
        record_data.get("section_height"),
    )


def export_aqueduct_summary(
    records,
    file_path,
):
    """
    附表2.3详细汇总导出的兼容入口。

    实际执行统一交给工程调查通用汇总导出器。
    """

    return export_engineering_summary(
        FORM_2_3,
        records=records,
        file_path=file_path,
    )


def _get_original_form_template_path():
    """
    获取附表2.3正式Excel模板路径。
    """

    return get_app_root() / "templates" / "excel" / "form_2_3_V1.xlsx"


def export_aqueduct_original_form(
    survey_record_id,
    file_path,
):
    """
    将一条附表2.3渡槽（座槽）
    调查记录填入正式原表模板。
    """

    # =========================================================
    # 1. 模板
    # =========================================================

    template_path = _get_original_form_template_path()

    if not template_path.exists():
        raise FileNotFoundError(
            "未找到附表2.3 Excel 模板。\n\n" f"应存在于：\n{template_path}"
        )

    # =========================================================
    # 2. 调查记录
    # =========================================================

    record = get_engineering_record(
            FORM_2_3,
            survey_record_id=survey_record_id,
        )

    if record is None:
        raise ValueError("没有找到需要导出的" "渡槽（座槽）调查记录。")

    asset = get_engineering_asset_detail(record["engineering_asset_id"])

    if asset is None:
        raise ValueError("没有找到该调查记录" "对应的工程对象。")

    inspection_results = get_inspection_results(survey_record_id)

    evaluation_map = {
        result["item_code"]: result["grade"] for result in inspection_results
    }

    record_data = record["record_data"] or {}

    # =========================================================
    # 3. 打开模板
    # =========================================================

    workbook = load_workbook(template_path)

    if "附表2.3" not in (workbook.sheetnames):
        workbook.close()

        raise ValueError("附表2.3 Excel模板中" "缺少工作表“附表2.3”。")

    worksheet = workbook["附表2.3"]

    # =========================================================
    # 4. 顶部归属和业务编号
    # =========================================================

    fill_original_form_ownership_header(
        worksheet,
        asset=asset,
        canal_id=record["canal_id"],
        business_code=record["business_code"],
    )

    # =========================================================
    # 5. 基本信息
    # =========================================================

    # 名称
    worksheet["B5"] = record["asset_name"] or ""

    # 桩号
    worksheet["H5"] = record_data.get("stake") or ""

    # 设计流量
    worksheet["J5"] = record_data.get("design_flow")

    # 建筑物等级
    worksheet["B6"] = record_data.get("structure_grade") or ""

    # 建成年月
    worksheet["D6"] = record_data.get("build_date") or ""

    # 加固改造年月
    worksheet["F6"] = record_data.get("renovation_date") or ""

    # 长度
    worksheet["H6"] = record_data.get("length")

    # 加大流量
    worksheet["J6"] = record_data.get("increased_flow")

    # 结构形式
    worksheet["B7"] = record_data.get("structure_form") or ""

    # 断面尺寸（宽*高）
    worksheet["D7"] = _format_section_size(record_data)

    # 槽身结构
    worksheet["F7"] = record_data.get("trough_body_structure") or ""

    # 槽壁厚度
    worksheet["H7"] = record_data.get("trough_wall_thickness")

    # 止水形式
    worksheet["J7"] = record_data.get("waterstop_form") or ""

    # 槽底高程
    worksheet["B8"] = record_data.get("trough_bottom_elevation")

    # 跨数
    worksheet["D8"] = record_data.get("span_count")

    # 下部支撑结构型式
    # F8:J8为合并单元格，
    # 只向左上角F8写值。
    worksheet["F8"] = record_data.get("lower_support_structure_form") or ""

    # =========================================================
    # 6. 12项分项评价
    # =========================================================
    #
    # 模板第10～21行，
    # 项目类别位于E列。
    # 与 FORM_2_3.evaluation_items
    # 顺序一一对应。
    # =========================================================

    for row_number, item in enumerate(
        FORM_2_3.evaluation_items,
        start=10,
    ):
        grade = evaluation_map.get(
            item["item_code"],
            "",
        )

        worksheet[f"E{row_number}"] = grade or ""

    # =========================================================
    # 7. 调查结论
    # =========================================================

    # 调查意见与建议
    # C22:H22为合并区域。
    worksheet["C22"] = record["survey_comment"] or ""

    # 工程状况类别
    worksheet["J22"] = record["overall_grade"] or ""

    # 调查时间
    worksheet["J23"] = record["survey_date"] or ""

    # B23 / D23 / F23 / H23
    # 为人工签字区域，程序保持空白。

    # =========================================================
    # 8. 打印设置
    # =========================================================

    worksheet.print_area = "A1:J26"

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1

    page_setup_properties = worksheet.sheet_properties.pageSetUpPr

    if page_setup_properties is not None:
        page_setup_properties.fitToPage = True

    worksheet.sheet_view.showGridLines = False

    # =========================================================
    # 9. 保存
    # =========================================================

    workbook.save(file_path)

    workbook.close()

    return {
        "file_path": file_path,
        "survey_record_id": (survey_record_id),
        "business_code": (record["business_code"]),
        "asset_name": (record["asset_name"]),
    }
