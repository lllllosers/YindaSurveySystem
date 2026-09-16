from openpyxl import (
    Workbook,
    load_workbook,
)

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

from openpyxl.utils import (
    get_column_letter,
)

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
)

from forms.engineering.form_2_5 import (
    FORM_2_5,
)

from forms.engineering.persistence import (
    get_engineering_record,
)

from services.engineering_summary_export import (
    export_engineering_summary,
)

from services.original_form_export_common import (
    fill_original_form_ownership_header,
)


def export_tunnel_summary(
    records,
    file_path,
):
    """
    详细汇总导出的兼容入口。

    实际执行统一交给工程调查通用汇总导出器。
    """

    return export_engineering_summary(
        FORM_2_5,
        records=records,
        file_path=file_path,
    )

def _get_original_form_template_path():
    """
    获取附表2.5正式Excel模板路径。
    """

    return get_app_root() / "templates" / "excel" / "form_2_5_V1.xlsx"


def _build_stake_range_text(
    start_stake,
    end_stake,
):
    start_text = str(start_stake or "").strip()

    end_text = str(end_stake or "").strip()

    if start_text and end_text:
        return f"{start_text}～" f"{end_text}"

    return start_text or end_text


def _build_section_size(
    width,
    height,
):
    """
    正式原表只有一个
    “尺寸(宽*高）”字段。

    程序内部仍保留 width / height
    两个字段，仅在导出层进行组合。
    """

    if width in (
        None,
        "",
    ) and height in (
        None,
        "",
    ):
        return ""

    width_text = "" if width is None else str(width)

    height_text = "" if height is None else str(height)

    if width_text and height_text:
        return f"{width_text}*" f"{height_text}"

    return width_text or height_text


def export_tunnel_original_form(
    survey_record_id,
    file_path,
):
    """
    将一条附表2.5隧洞调查记录
    填入正式原表模板。
    """

    # =========================================================
    # 1. 模板
    # =========================================================

    template_path = _get_original_form_template_path()

    if not template_path.exists():
        raise FileNotFoundError(
            "未找到附表2.5 Excel 模板。" "\n\n" f"应存在于：\n" f"{template_path}"
        )

    # =========================================================
    # 2. 调查记录
    # =========================================================

    record = get_engineering_record(
            FORM_2_5,
            survey_record_id=survey_record_id,
        )

    if record is None:
        raise ValueError("没有找到需要导出的" "隧洞调查记录。")

    asset = get_engineering_asset_detail(record["engineering_asset_id"])

    if asset is None:
        raise ValueError("没有找到该调查记录" "对应的工程对象。")

    inspection_results = get_inspection_results(survey_record_id)

    evaluation_map = {
        result["item_code"]: result["grade"] for result in inspection_results
    }

    record_data = record["record_data"] or {}

    # =========================================================
    # 3. 打开模板
    # =========================================================

    workbook = load_workbook(template_path)

    if "附表2.5" not in workbook.sheetnames:
        workbook.close()

        raise ValueError("附表2.5 Excel模板中" "缺少工作表“附表2.5”。")

    worksheet = workbook["附表2.5"]

    # =========================================================
    # 4. 顶部归属和业务编号
    # =========================================================

    fill_original_form_ownership_header(
        worksheet,
        asset=asset,
        canal_id=record["canal_id"],
        business_code=record["business_code"],
    )

    # =========================================================
    # 5. 基本信息
    # =========================================================

    worksheet["B5"] = record["asset_name"] or ""

    worksheet["F5"] = _build_stake_range_text(
        record["start_stake_text"],
        record["end_stake_text"],
    )

    worksheet["J5"] = record_data.get("design_flow")

    worksheet["B6"] = record_data.get("structure_grade") or ""

    worksheet["D6"] = record_data.get("build_date") or ""

    worksheet["F6"] = record_data.get("renovation_date") or ""

    worksheet["H6"] = record_data.get("length")

    worksheet["J6"] = record_data.get("increased_flow")

    worksheet["B7"] = record_data.get("lining_form") or ""

    worksheet["D7"] = record_data.get("lining_thickness")

    worksheet["F7"] = record_data.get("concrete_strength") or ""

    worksheet["H7"] = record_data.get("inlet_outlet_bottom_elevation") or ""

    worksheet["J7"] = record_data.get("longitudinal_slope")

    worksheet["B8"] = record_data.get("section_form") or ""

    worksheet["E8"] = _build_section_size(
        record_data.get("section_width"),
        record_data.get("section_height"),
    )

    worksheet["J8"] = record_data.get("cover_thickness")

    # =========================================================
    # 6. 13项分项评价
    # =========================================================

    for (
        row_number,
        item,
    ) in enumerate(
        FORM_2_5.evaluation_items,
        start=10,
    ):
        grade = evaluation_map.get(
            item["item_code"],
            "",
        )

        worksheet[f"E{row_number}"] = grade or ""

    # =========================================================
    # 7. 调查结论
    # =========================================================

    worksheet["C23"] = record["survey_comment"] or ""

    worksheet["J23"] = record["overall_grade"] or ""

    # 签字栏不由软件填写。
    worksheet["J24"] = record["survey_date"] or ""

    # =========================================================
    # 8. 打印设置
    # =========================================================

    worksheet.print_area = "A1:J25"

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1

    page_setup_properties = worksheet.sheet_properties.pageSetUpPr

    if page_setup_properties is not None:
        page_setup_properties.fitToPage = True

    worksheet.sheet_view.showGridLines = False

    workbook.save(file_path)

    workbook.close()

    return {
        "file_path": file_path,
        "survey_record_id": (survey_record_id),
    }
