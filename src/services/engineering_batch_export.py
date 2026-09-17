from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import shutil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from database import (
    get_app_root,
    get_canal_units,
    get_departments,
    get_engineering_survey_query_records,
    get_water_offices,
)
from forms.engineering.registry import (
    get_engineering_asset_type_display_name,
    get_engineering_form_definitions,
)
from services.engineering_original_form_export import (
    export_engineering_original_form,
)
from services.engineering_summary_export import (
    export_engineering_summary,
)
from services.survey_scope import (
    SurveyScope,
    filter_records_by_scope,
    resolve_survey_scope,
)


_INVALID_FILENAME_CHARS = '\\/:*?"<>|'


@dataclass(frozen=True)
class BatchExportRequest:
    scope: SurveyScope
    output_root: Path
    create_zip: bool = False

    def __post_init__(self):
        if not isinstance(self.scope, SurveyScope):
            raise TypeError("scope 必须是 SurveyScope。")

        object.__setattr__(
            self,
            "output_root",
            Path(self.output_root),
        )


@dataclass(frozen=True)
class FormExportGroup:
    definition: object
    records: tuple


@dataclass(frozen=True)
class BatchExportPlan:
    request: BatchExportRequest
    records: tuple
    groups: tuple[FormExportGroup, ...]


@dataclass(frozen=True)
class BatchExportResult:
    output_root: Path
    manifest_path: Path
    archive_path: Path | None
    total_records: int
    summary_success_count: int
    original_success_count: int
    errors: tuple[str, ...]

    @property
    def completed(self):
        return not self.errors


def sanitize_filename(
    value,
    *,
    fallback="未命名",
):
    """生成适合 Windows 成果目录使用的文件名片段。"""

    text = str(value or "").strip()

    for char in _INVALID_FILENAME_CHARS:
        text = text.replace(char, "_")

    text = " ".join(text.split()).rstrip(". ")

    if not text:
        text = fallback

    reserved_names = {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *{f"COM{index}" for index in range(1, 10)},
        *{f"LPT{index}" for index in range(1, 10)},
    }

    if text.upper() in reserved_names:
        text = f"_{text}"

    return text


def _form_output_name(definition):
    short_name = (
        get_engineering_asset_type_display_name(
            definition.asset_type
        )
        or definition.form_name
    )

    return sanitize_filename(
        f"附表{definition.form_number}_{short_name}"
    )


def _validate_export_support(definition):
    if definition.summary_export_definition is None:
        raise ValueError(
            f"{definition.form_code} "
            "尚未配置详细汇总导出定义。"
        )

    original_definition = (
        definition.original_form_export_definition
    )

    if original_definition is None:
        raise ValueError(
            f"{definition.form_code} "
            "尚未配置正式原表导出定义。"
        )

    template_path = (
        get_app_root()
        / "templates"
        / "excel"
        / original_definition.template_filename
    )

    if not template_path.exists():
        raise FileNotFoundError(
            "正式原表模板不存在："
            f"{template_path}"
        )


def load_scope_records(scope):
    """
    从当前数据库读取 SurveyScope 对应的工程调查记录。

    组织和渠系均按“所选根节点 + 下级节点”语义解析。
    正式成果默认由 SurveyScope 限制为 completed。
    """

    organization_rows = (
        list(get_departments())
        + list(get_water_offices())
    )

    canal_rows = list(get_canal_units())

    resolved_scope = resolve_survey_scope(
        scope,
        organization_rows=organization_rows,
        canal_rows=canal_rows,
    )

    records = get_engineering_survey_query_records(
        project_id=scope.project_id,
        survey_batch_id=scope.survey_batch_id,
    )

    return filter_records_by_scope(
        records,
        resolved_scope,
    )


def build_batch_export_plan(request):
    """
    读取、筛选并预检附表2正式成果导出计划。

    计划构建阶段不写文件。
    """

    if not isinstance(request, BatchExportRequest):
        raise TypeError(
            "request 必须是 BatchExportRequest。"
        )

    scope = request.scope

    registered_definitions = (
        get_engineering_form_definitions()
    )

    definition_by_code = {
        definition.form_code: definition
        for definition in registered_definitions
    }

    if scope.form_codes:
        unknown_form_codes = tuple(
            form_code
            for form_code in scope.form_codes
            if form_code not in definition_by_code
        )

        if unknown_form_codes:
            raise ValueError(
                "调查范围包含未注册表单："
                + "、".join(unknown_form_codes)
            )

    records = tuple(load_scope_records(scope))

    if not records:
        raise ValueError(
            "当前调查范围没有可生成正式成果的"
            "已完成工程调查记录。"
        )

    records_by_form = {}

    for record in records:
        form_code = record["form_code"]

        definition = definition_by_code.get(
            form_code
        )

        if definition is None:
            raise ValueError(
                "调查记录引用了未注册工程调查表："
                f"{form_code}"
            )

        records_by_form.setdefault(
            form_code,
            [],
        ).append(record)

    groups = []

    for definition in registered_definitions:
        group_records = records_by_form.get(
            definition.form_code
        )

        if not group_records:
            continue

        _validate_export_support(definition)

        groups.append(
            FormExportGroup(
                definition=definition,
                records=tuple(group_records),
            )
        )

    return BatchExportPlan(
        request=request,
        records=records,
        groups=tuple(groups),
    )


def _ensure_empty_output_root(output_root):
    if output_root.exists():
        if not output_root.is_dir():
            raise ValueError(
                "成果导出目标必须是目录。"
            )

        if any(output_root.iterdir()):
            raise ValueError(
                "成果导出目标目录不是空目录。"
            )
    else:
        output_root.mkdir(
            parents=True,
            exist_ok=False,
        )


def _reserve_unique_path(
    directory,
    filename,
    *,
    survey_record_id,
):
    path = directory / filename

    if not path.exists():
        return path

    return directory / (
        f"{path.stem}__记录"
        f"{survey_record_id}"
        f"{path.suffix}"
    )


def _write_manifest(
    *,
    plan,
    manifest_path,
    summary_entries,
    record_entries,
    errors,
):
    workbook = Workbook()

    overview = workbook.active

    if overview is None:
        overview = workbook.create_sheet(
            "导出概览"
        )
    else:
        overview.title = "导出概览"

    scope = plan.request.scope

    overview_rows = [
        ("项目ID", scope.project_id),
        ("调查批次ID", scope.survey_batch_id),
        ("正式成果记录数", len(plan.records)),
        ("涉及调查表数", len(plan.groups)),
        (
            "导出结论",
            "导出完成"
            if not errors
            else "导出未完成",
        ),
        ("失败项数", len(errors)),
    ]

    for key, value in overview_rows:
        overview.append([key, value])

    overview.column_dimensions["A"].width = 22
    overview.column_dimensions["B"].width = 56

    for cell in overview["A"]:
        cell.font = Font(bold=True)

    summary_sheet = workbook.create_sheet(
        "汇总文件"
    )

    summary_sheet.append(
        [
            "调查表",
            "记录数",
            "导出状态",
            "相对路径",
            "错误信息",
        ]
    )

    for entry in summary_entries:
        summary_sheet.append(
            [
                entry["form_display_name"],
                entry["record_count"],
                entry["status"],
                entry["relative_path"],
                entry["error"],
            ]
        )

    record_sheet = workbook.create_sheet(
        "记录清单"
    )

    record_sheet.append(
        [
            "序号",
            "调查表",
            "业务编号",
            "工程名称",
            "基层处",
            "水管所",
            "渠系",
            "工程位置",
            "工程状况类别",
            "调查时间",
            "正式原表状态",
            "正式原表相对路径",
            "错误信息",
        ]
    )

    for index, entry in enumerate(
        record_entries,
        start=1,
    ):
        record_sheet.append(
            [
                index,
                entry["form_display_name"],
                entry["business_code"],
                entry["asset_name"],
                entry["department_name"],
                entry["office_name"],
                entry["canal_name"],
                entry["engineering_position"],
                entry["overall_grade"],
                entry["survey_date"],
                entry["status"],
                entry["relative_path"],
                entry["error"],
            ]
        )

    for sheet in (
        summary_sheet,
        record_sheet,
    ):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            sheet.dimensions
        )
        sheet.sheet_view.showGridLines = False

        for row in sheet.iter_rows(
            min_row=2,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

    summary_widths = [
        34,
        10,
        14,
        55,
        60,
    ]

    for index, width in enumerate(
        summary_widths,
        start=1,
    ):
        summary_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    record_widths = [
        8,
        34,
        24,
        24,
        16,
        16,
        18,
        22,
        14,
        14,
        16,
        65,
        60,
    ]

    for index, width in enumerate(
        record_widths,
        start=1,
    ):
        record_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    workbook.save(manifest_path)
    workbook.close()


def execute_batch_export(plan):
    """
    执行附表2批量正式成果导出。

    失败策略：
    - 单张汇总失败，不阻断其他表单；
    - 单条正式原表失败，不阻断其他记录；
    - 最终始终尝试生成成果清单；
    - 只要存在失败项，就写“导出未完成.txt”；
    - 只有全部成功时才允许生成 ZIP。
    """

    if not isinstance(plan, BatchExportPlan):
        raise TypeError(
            "plan 必须是 BatchExportPlan。"
        )

    output_root = plan.request.output_root

    _ensure_empty_output_root(
        output_root
    )

    summary_root = (
        output_root / "01_详细汇总"
    )

    original_root = (
        output_root / "02_正式调查表"
    )

    summary_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    original_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    errors = []
    summary_entries = []
    record_entries = []

    summary_success_count = 0
    original_success_count = 0

    for group in plan.groups:
        definition = group.definition

        form_output_name = _form_output_name(
            definition
        )

        form_display_name = (
            f"附表{definition.form_number} "
            f"{definition.form_name}"
        )

        summary_path = (
            summary_root
            / f"{form_output_name}_汇总.xlsx"
        )

        summary_relative_path = (
            summary_path.relative_to(
                output_root
            )
        )

        try:
            export_engineering_summary(
                definition,
                records=group.records,
                file_path=summary_path,
            )

            summary_status = "成功"
            summary_error = ""
            summary_success_count += 1

        except Exception as error:
            summary_status = "失败"
            summary_error = str(error)

            errors.append(
                f"{form_display_name}"
                " 详细汇总失败："
                f"{error}"
            )

        summary_entries.append(
            {
                "form_display_name": (
                    form_display_name
                ),
                "record_count": len(
                    group.records
                ),
                "status": summary_status,
                "relative_path": str(
                    summary_relative_path
                ),
                "error": summary_error,
            }
        )

        form_original_root = (
            original_root
            / form_output_name
        )

        form_original_root.mkdir(
            parents=True,
            exist_ok=True,
        )

        for record in group.records:
            survey_record_id = int(
                record["survey_record_id"]
            )

            file_stem = sanitize_filename(
                (
                    f"{record['canal_name']}_"
                    f"{record['business_code']}_"
                    f"{record['asset_name']}"
                ),
                fallback=(
                    f"记录{survey_record_id}"
                ),
            )

            original_path = _reserve_unique_path(
                form_original_root,
                f"{file_stem}.xlsx",
                survey_record_id=(
                    survey_record_id
                ),
            )

            original_relative_path = (
                original_path.relative_to(
                    output_root
                )
            )

            try:
                export_engineering_original_form(
                    definition,
                    survey_record_id=(
                        survey_record_id
                    ),
                    file_path=original_path,
                )

                original_status = "成功"
                original_error = ""
                original_success_count += 1

            except Exception as error:
                original_status = "失败"
                original_error = str(error)

                errors.append(
                    f"{form_display_name} "
                    f"记录{survey_record_id}"
                    " 正式原表失败："
                    f"{error}"
                )

            record_entries.append(
                {
                    "form_display_name": (
                        form_display_name
                    ),
                    "business_code": (
                        record["business_code"]
                    ),
                    "asset_name": (
                        record["asset_name"]
                    ),
                    "department_name": (
                        record["department_name"]
                    ),
                    "office_name": (
                        record["office_name"]
                    ),
                    "canal_name": (
                        record["canal_name"]
                    ),
                    "engineering_position": (
                        record[
                            "engineering_position"
                        ]
                    ),
                    "overall_grade": (
                        record["overall_grade"]
                        or ""
                    ),
                    "survey_date": (
                        record["survey_date"]
                    ),
                    "status": original_status,
                    "relative_path": str(
                        original_relative_path
                    ),
                    "error": original_error,
                }
            )

    manifest_path = (
        output_root / "00_成果清单.xlsx"
    )

    _write_manifest(
        plan=plan,
        manifest_path=manifest_path,
        summary_entries=summary_entries,
        record_entries=record_entries,
        errors=errors,
    )

    failure_marker = (
        output_root / "导出未完成.txt"
    )

    archive_path = None

    if errors:
        failure_marker.write_text(
            (
                "本次成果导出存在失败项，"
                "不能作为完整正式成果。\n\n"
                + "\n".join(
                    f"{index}. {message}"
                    for index, message
                    in enumerate(
                        errors,
                        start=1,
                    )
                )
                + "\n"
            ),
            encoding="utf-8",
        )

    elif plan.request.create_zip:
        archive_path = Path(
            shutil.make_archive(
                str(output_root),
                "zip",
                root_dir=output_root.parent,
                base_dir=output_root.name,
            )
        )

    return BatchExportResult(
        output_root=output_root,
        manifest_path=manifest_path,
        archive_path=archive_path,
        total_records=len(plan.records),
        summary_success_count=(
            summary_success_count
        ),
        original_success_count=(
            original_success_count
        ),
        errors=tuple(errors),
    )


def export_engineering_batch(request):
    """
    一步完成“计划构建 + 批量正式成果导出”。

    UI 后续只需要构造 BatchExportRequest，
    不直接操作14张表各自的 Exporter。
    """

    plan = build_batch_export_plan(
        request
    )

    return execute_batch_export(
        plan
    )
