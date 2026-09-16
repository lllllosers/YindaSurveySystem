from openpyxl import (
    Workbook,
    load_workbook,
)

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

from openpyxl.utils import (
    get_column_letter,
)

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
    get_point_engineering_record,
)

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
    get_point_engineering_record,
)

from forms.engineering.form_2_6 import (
    FORM_2_6,
)

from services.culvert_evaluation import (
    CULVERT_EVALUATION_ITEMS,
)

from services.culvert_evaluation import (
    CULVERT_EVALUATION_ITEMS,
)

from services.original_form_export_common import (
    fill_original_form_ownership_header,
)

# =============================================================
# 详细汇总导出
# =============================================================


def export_culvert_summary(
    records,
    file_path,
):
    """
    将当前筛选后的附表2.6
    涵洞（暗涵）记录
    导出为详细数据汇总 Excel。

    一条 SurveyRecord 对应一行。
    """

    if not records:
        raise ValueError("当前没有可导出的调查记录。")

    workbook = Workbook()

    worksheet = workbook.active

    if worksheet is None:
        worksheet = workbook.create_sheet("涵洞（暗涵）调查汇总")
    else:
        worksheet.title = "涵洞（暗涵）调查汇总"

    # =========================================================
    # 表头
    # =========================================================

    basic_headers = [
        "序号",
        "业务编号",
        "名称",
        "基层处",
        "水管所",
        "渠系",
        "桩号",
        "设计流量（m³/s）",
        "建筑物等级",
        "建成年月",
        "加固改造年月",
        "长度",
        "加大流量（m³/s）",
        "结构形式",
        "主构建筑材料",
        "混凝土强度",
        "钢筋保护层厚度",
        "覆土厚度（m）",
        "渠宽（m）",
        "渠深（m）",
        "渠底高程（m）",
    ]

    evaluation_headers = [
        (f"{item['category']}" f"-{item['item_name']}")
        for item in CULVERT_EVALUATION_ITEMS
    ]

    conclusion_headers = [
        "工程状况类别",
        "调查时间",
        "调查意见与建议",
        "状态",
        "修改时间",
    ]

    headers = basic_headers + evaluation_headers + conclusion_headers

    worksheet.append(headers)

    # =========================================================
    # 数据
    # =========================================================

    exported_count = 0

    for index, summary_record in enumerate(
        records,
        start=1,
    ):
        survey_record_id = summary_record["survey_record_id"]

        record = get_point_engineering_record(
            survey_record_id=(survey_record_id),
            form_code=FORM_2_6.form_code,
        )

        if record is None:
            raise ValueError(
                "导出过程中发现" "附表2.6调查记录不存在：" f"{survey_record_id}"
            )

        record_data = record["record_data"] or {}

        inspection_results = get_inspection_results(survey_record_id)

        evaluation_map = {
            result["item_code"]: result["grade"] for result in inspection_results
        }

        status_text = {
            "draft": "草稿",
            "completed": "录入完成",
        }.get(
            record["record_status"],
            record["record_status"],
        )

        row = [
            index,
            record["business_code"],
            record["asset_name"],
            summary_record["department_name"],
            summary_record["office_name"],
            summary_record["canal_name"],
            record_data.get("stake"),
            record_data.get("design_flow"),
            record_data.get("structure_grade"),
            record_data.get("build_date"),
            record_data.get("renovation_date"),
            record_data.get("length"),
            record_data.get("increased_flow"),
            record_data.get("structure_form"),
            record_data.get("main_structure_material"),
            record_data.get("concrete_strength"),
            record_data.get("cover_thickness"),
            record_data.get("soil_cover_thickness"),
            record_data.get("channel_width"),
            record_data.get("channel_depth"),
            record_data.get("channel_bottom_elevation"),
        ]

        for item in CULVERT_EVALUATION_ITEMS:
            row.append(
                evaluation_map.get(
                    item["item_code"],
                    "",
                )
            )

        row.extend(
            [
                record["overall_grade"],
                record["survey_date"],
                record["survey_comment"],
                status_text,
                summary_record["updated_at"],
            ]
        )

        worksheet.append(row)

        exported_count += 1

    # =========================================================
    # 样式
    # =========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="D9DEE3",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = border

    worksheet.row_dimensions[1].height = 36

    for row_cells in worksheet.iter_rows(
        min_row=2,
        max_row=(worksheet.max_row),
    ):
        for cell in row_cells:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            cell.border = border

    # =========================================================
    # 居中列
    # =========================================================

    center_columns = {
        1,  # 序号
        7,  # 桩号
        8,  # 设计流量
        9,  # 建筑物等级
        10,  # 建成年月
        11,  # 加固改造年月
        12,  # 长度
        13,  # 加大流量
        17,  # 钢筋保护层厚度
        18,  # 覆土厚度
        19,  # 渠宽
        20,  # 渠深
        21,  # 渠底高程
    }

    evaluation_start_column = 22

    evaluation_end_column = evaluation_start_column + len(CULVERT_EVALUATION_ITEMS) - 1

    center_columns.update(
        range(
            evaluation_start_column,
            evaluation_end_column + 1,
        )
    )

    overall_grade_column = evaluation_end_column + 1

    survey_date_column = evaluation_end_column + 2

    comment_column = evaluation_end_column + 3

    status_column = evaluation_end_column + 4

    updated_at_column = evaluation_end_column + 5

    center_columns.update(
        {
            overall_grade_column,
            survey_date_column,
            status_column,
        }
    )

    for column_index in center_columns:
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # =========================================================
    # 列宽
    # =========================================================

    widths = {
        1: 8,
        2: 22,
        3: 20,
        4: 16,
        5: 16,
        6: 18,
        7: 14,
        8: 16,
        9: 14,
        10: 14,
        11: 16,
        12: 12,
        13: 16,
        14: 18,
        15: 18,
        16: 16,
        17: 18,
        18: 16,
        19: 12,
        20: 12,
        21: 16,
    }

    for column_index in range(
        evaluation_start_column,
        evaluation_end_column + 1,
    ):
        widths[column_index] = 20

    widths[overall_grade_column] = 14

    widths[survey_date_column] = 14

    widths[comment_column] = 36

    widths[status_column] = 12

    widths[updated_at_column] = 20

    for (
        column_index,
        width,
    ) in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    # =========================================================
    # 使用体验
    # =========================================================

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    workbook.save(file_path)

    return {
        "file_path": file_path,
        "exported_count": (exported_count),
    }


# =============================================================
# 正式原表
# =============================================================


def _get_original_form_template_path():
    """
    获取附表2.6正式Excel模板路径。
    """

    return get_app_root() / "templates" / "excel" / "form_2_6_V1.xlsx"


def export_culvert_original_form(
    survey_record_id,
    file_path,
):
    """
    将一条附表2.6涵洞（暗涵）
    调查记录填入正式原表模板。
    """

    # =========================================================
    # 1. 模板
    # =========================================================

    template_path = _get_original_form_template_path()

    if not template_path.exists():
        raise FileNotFoundError(
            "未找到附表2.6 Excel 模板。" "\n\n" f"应存在于：\n" f"{template_path}"
        )

    # =========================================================
    # 2. 调查记录
    # =========================================================

    record = get_point_engineering_record(
        survey_record_id=(survey_record_id),
        form_code=FORM_2_6.form_code,
    )

    if record is None:
        raise ValueError("没有找到需要导出的" "涵洞（暗涵）调查记录。")

    asset = get_engineering_asset_detail(record["engineering_asset_id"])

    if asset is None:
        raise ValueError("没有找到该调查记录" "对应的工程对象。")

    inspection_results = get_inspection_results(survey_record_id)

    evaluation_map = {
        result["item_code"]: result["grade"] for result in inspection_results
    }

    record_data = record["record_data"] or {}

    # =========================================================
    # 3. 打开模板
    # =========================================================

    workbook = load_workbook(template_path)

    if "附表2.6" not in workbook.sheetnames:
        workbook.close()

        raise ValueError("附表2.6 Excel模板中" "缺少工作表“附表2.6”。")

    worksheet = workbook["附表2.6"]

    # =========================================================
    # 4. 顶部归属和业务编号
    # =========================================================

    fill_original_form_ownership_header(
        worksheet,
        asset=asset,
        canal_id=record["canal_id"],
        business_code=(record["business_code"]),
    )

    # =========================================================
    # 5. 基本信息
    # =========================================================

    # 第一行
    worksheet["B5"] = record["asset_name"] or ""

    worksheet["H5"] = record_data.get("stake") or ""

    worksheet["J5"] = record_data.get("design_flow")

    # 第二行
    worksheet["B6"] = record_data.get("structure_grade") or ""

    worksheet["D6"] = record_data.get("build_date") or ""

    worksheet["F6"] = record_data.get("renovation_date") or ""

    worksheet["H6"] = record_data.get("length")

    worksheet["J6"] = record_data.get("increased_flow")

    # 第三行
    worksheet["B7"] = record_data.get("structure_form") or ""

    worksheet["D7"] = record_data.get("main_structure_material") or ""

    worksheet["F7"] = record_data.get("concrete_strength") or ""

    worksheet["H7"] = record_data.get("cover_thickness")

    worksheet["J7"] = record_data.get("soil_cover_thickness")

    # 第四行
    worksheet["B8"] = record_data.get("channel_width")

    worksheet["D8"] = record_data.get("channel_depth")

    worksheet["F8"] = record_data.get("channel_bottom_elevation")

    # =========================================================
    # 6. 11项分项评价
    # =========================================================

    for (
        row_number,
        item,
    ) in enumerate(
        CULVERT_EVALUATION_ITEMS,
        start=10,
    ):
        grade = evaluation_map.get(
            item["item_code"],
            "",
        )

        worksheet[f"E{row_number}"] = grade or ""

    # =========================================================
    # 7. 调查结论
    # =========================================================

    worksheet["C21"] = record["survey_comment"] or ""

    worksheet["J21"] = record["overall_grade"] or ""

    # 签字栏不由软件填写。
    worksheet["J22"] = record["survey_date"] or ""

    # =========================================================
    # 8. 打印设置
    # =========================================================

    worksheet.print_area = "A1:J23"

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1

    page_setup_properties = worksheet.sheet_properties.pageSetUpPr

    if page_setup_properties is not None:
        page_setup_properties.fitToPage = True

    worksheet.sheet_view.showGridLines = False

    workbook.save(file_path)

    workbook.close()

    return {
        "file_path": file_path,
        "survey_record_id": (survey_record_id),
    }
