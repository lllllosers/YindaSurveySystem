from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from database import get_inspection_results

from forms.engineering.formatters import (
    format_record_status,
)
from forms.engineering.models import (
    EngineeringFormDefinition,
)
from forms.engineering.persistence import (
    get_engineering_record,
)


_CONCLUSION_HEADERS = (
    "工程状况类别",
    "调查时间",
    "调查意见与建议",
    "状态",
    "修改时间",
)

_CONCLUSION_WIDTHS = (
    14,
    14,
    36,
    12,
    20,
)

_CONCLUSION_ALIGNMENTS = (
    "center",
    "center",
    "left",
    "center",
    "left",
)


def _resolve_binding(
    *,
    summary_record,
    record,
    record_data,
    binding,
):
    """解析 SummaryColumnDefinition 的声明式值绑定。"""

    if binding.source == "query_record":
        values = [
            summary_record[key]
            for key in binding.keys
        ]

    elif binding.source == "record":
        values = [
            record[key]
            for key in binding.keys
        ]

    elif binding.source == "record_data":
        values = [
            record_data.get(key)
            for key in binding.keys
        ]

    else:
        raise ValueError(
            "详细汇总不支持的数据源："
            f"{binding.source}"
        )

    if binding.formatter is not None:
        return binding.formatter(*values)

    return values[0]


def _apply_data_alignment(
    cell,
    alignment,
    border,
):
    cell.alignment = Alignment(
        horizontal=alignment,
        vertical="center",
        wrap_text=True,
    )
    cell.border = border


def export_engineering_summary(
    definition: EngineeringFormDefinition,
    *,
    records,
    file_path,
):
    """
    按 EngineeringFormDefinition.summary_export_definition
    导出一张附表2的详细数据汇总。

    “序号”、分项评价和调查结论属于所有工程调查表
    的公共语义，由本函数统一生成；各表只声明自己的
    基础信息列。
    """

    summary_definition = (
        definition.summary_export_definition
    )

    if summary_definition is None:
        raise ValueError(
            f"{definition.form_code} "
            "尚未配置 SummaryExportDefinition。"
        )

    if not records:
        raise ValueError(
            "当前没有可导出的调查记录。"
        )

    workbook = Workbook()

    worksheet = workbook.active

    if worksheet is None:
        worksheet = workbook.create_sheet(
            summary_definition.sheet_name
        )
    else:
        worksheet.title = (
            summary_definition.sheet_name
        )

    evaluation_headers = [
        (
            f"{item['category']}"
            f"-{item['item_name']}"
        )
        for item in definition.evaluation_items
    ]

    headers = (
        ["序号"]
        + [
            column.header
            for column
            in summary_definition.columns
        ]
        + evaluation_headers
        + list(_CONCLUSION_HEADERS)
    )

    worksheet.append(headers)

    exported_count = 0

    for index, summary_record in enumerate(
        records,
        start=1,
    ):
        survey_record_id = int(
            summary_record["survey_record_id"]
        )

        record = get_engineering_record(
            definition,
            survey_record_id=survey_record_id,
        )

        if record is None:
            raise ValueError(
                "导出过程中发现调查记录不存在："
                f"{survey_record_id}"
            )

        record_data = (
            record["record_data"]
            or {}
        )

        inspection_results = (
            get_inspection_results(
                survey_record_id
            )
        )

        evaluation_map = {
            result["item_code"]:
            result["grade"]
            for result
            in inspection_results
        }

        row = [index]

        row.extend(
            _resolve_binding(
                summary_record=summary_record,
                record=record,
                record_data=record_data,
                binding=column.binding,
            )
            for column
            in summary_definition.columns
        )

        row.extend(
            evaluation_map.get(
                item["item_code"],
                "",
            )
            for item
            in definition.evaluation_items
        )

        row.extend(
            [
                record["overall_grade"],
                record["survey_date"],
                record["survey_comment"],
                format_record_status(
                    record["record_status"]
                ),
                summary_record["updated_at"],
            ]
        )

        worksheet.append(row)

        exported_count += 1

    # =========================================================
    # 公共样式
    # =========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="D9DEE3",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    worksheet.row_dimensions[1].height = (
        summary_definition.header_height
    )

    # =========================================================
    # 列位置全部由定义长度推导，不保存硬编码索引
    # =========================================================

    basic_start_column = 2

    evaluation_start_column = (
        basic_start_column
        + len(summary_definition.columns)
    )

    evaluation_end_column = (
        evaluation_start_column
        + len(definition.evaluation_items)
        - 1
    )

    conclusion_start_column = (
        evaluation_end_column + 1
    )

    # 序号
    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        _apply_data_alignment(
            worksheet.cell(
                row=row_index,
                column=1,
            ),
            "center",
            border,
        )

    # 各表自己的基础信息列
    for offset, column_definition in enumerate(
        summary_definition.columns,
        start=basic_start_column,
    ):
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            _apply_data_alignment(
                worksheet.cell(
                    row=row_index,
                    column=offset,
                ),
                column_definition.alignment,
                border,
            )

    # 分项评价
    for column_index in range(
        evaluation_start_column,
        evaluation_end_column + 1,
    ):
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            _apply_data_alignment(
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ),
                "center",
                border,
            )

    # 调查结论
    for offset, alignment in enumerate(
        _CONCLUSION_ALIGNMENTS,
        start=conclusion_start_column,
    ):
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            _apply_data_alignment(
                worksheet.cell(
                    row=row_index,
                    column=offset,
                ),
                alignment,
                border,
            )

    # =========================================================
    # 列宽
    # =========================================================

    worksheet.column_dimensions["A"].width = 8

    for column_index, column_definition in enumerate(
        summary_definition.columns,
        start=basic_start_column,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = column_definition.width

    for column_index in range(
        evaluation_start_column,
        evaluation_end_column + 1,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = (
            summary_definition
            .evaluation_column_width
        )

    for column_index, width in enumerate(
        _CONCLUSION_WIDTHS,
        start=conclusion_start_column,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    # =========================================================
    # Excel 使用体验
    # =========================================================

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )
    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = (
        "landscape"
    )
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    workbook.save(file_path)

    return {
        "file_path": file_path,
        "exported_count": exported_count,
    }
