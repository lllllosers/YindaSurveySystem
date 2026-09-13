from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import (
    get_column_letter,
)


def export_common_query_summary(
    records,
    file_path,
):
    """
    将跨调查表查询结果导出为
    公共字段汇总 Excel。

    该导出只包含所有工程调查表共有的信息。
    """

    if not records:
        raise ValueError("当前没有可导出的查询结果。")

    workbook = Workbook()

    worksheet = workbook.active

    if worksheet is None:
        worksheet = workbook.create_sheet("统一数据查询")
    else:
        worksheet.title = "统一数据查询"

    headers = [
        "序号",
        "调查表",
        "调查批次",
        "业务编号",
        "工程名称",
        "基层处",
        "水管所",
        "渠系",
        "工程位置",
        "工程状况类别",
        "调查时间",
        "状态",
        "修改时间",
    ]

    worksheet.append(headers)

    for index, record in enumerate(
        records,
        start=1,
    ):
        status_text = {
            "draft": "草稿",
            "completed": "录入完成",
        }.get(
            record["record_status"],
            record["record_status"],
        )

        worksheet.append(
            [
                index,
                record["form_display_name"],
                record["batch_name"],
                record["business_code"],
                record["asset_name"],
                record["department_name"],
                record["office_name"],
                record["canal_name"],
                record["engineering_position"],
                record["overall_grade"],
                record["survey_date"],
                status_text,
                record["updated_at"],
            ]
        )

    # =========================================================
    # 样式
    # =========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="D9DEE3",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = border

    worksheet.row_dimensions[1].height = 36

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            cell.border = border

    center_columns = {
        1,
        9,
        10,
        11,
        12,
    }

    for column_index in center_columns:
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    widths = {
        1: 8,
        2: 34,
        3: 24,
        4: 22,
        5: 22,
        6: 16,
        7: 16,
        8: 18,
        9: 24,
        10: 14,
        11: 14,
        12: 12,
        13: 20,
    }

    for column_index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    workbook.save(file_path)

    return {
        "file_path": file_path,
        "exported_count": (len(records)),
    }
