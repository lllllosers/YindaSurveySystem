from openpyxl import load_workbook

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
)

from forms.engineering.models import (
    EngineeringFormDefinition,
)
from forms.engineering.persistence import (
    get_engineering_record,
)

from services.original_form_export_common import (
    fill_original_form_ownership_header,
)


def _resolve_binding(
    *,
    record,
    record_data,
    binding,
):
    """
    解析正式原表字段绑定。

    OriginalFormExportDefinition 仅允许
    record / record_data 两种来源。
    """

    if binding.source == "record":
        values = [
            record[key]
            for key in binding.keys
        ]

    elif binding.source == "record_data":
        values = [
            record_data.get(key)
            for key in binding.keys
        ]

    else:
        raise ValueError(
            "正式原表不支持的数据源："
            f"{binding.source}"
        )

    if binding.formatter is not None:
        return binding.formatter(*values)

    return values[0]


def _resolve_ownership_canal_id(
    record,
):
    """
    point 记录历史上暴露 canal_id，
    range 记录历史上暴露 canal_unit_id。

    正式原表顶部归属属于通用能力，
    在执行层兼容两种既有 record 合同，
    不把该差异泄漏到每张表的定义中。
    """

    for key in (
        "canal_id",
        "canal_unit_id",
    ):
        try:
            value = record[key]
        except (
            KeyError,
            IndexError,
            TypeError,
        ):
            continue

        if value is not None:
            return value

    raise ValueError(
        "调查记录缺少渠系归属信息。"
    )


def _apply_print_settings(
    worksheet,
    settings,
):
    worksheet.print_area = (
        settings.print_area
    )

    worksheet.page_setup.orientation = (
        settings.orientation
    )

    paper_sizes = {
        "A4": worksheet.PAPERSIZE_A4,
    }

    try:
        worksheet.page_setup.paperSize = (
            paper_sizes[settings.paper_size]
        )
    except KeyError as error:
        raise ValueError(
            "通用正式原表导出器暂不支持纸张："
            f"{settings.paper_size}"
        ) from error

    worksheet.page_setup.fitToWidth = (
        settings.fit_to_width
    )
    worksheet.page_setup.fitToHeight = (
        settings.fit_to_height
    )

    page_setup_properties = (
        worksheet
        .sheet_properties
        .pageSetUpPr
    )

    if page_setup_properties is not None:
        page_setup_properties.fitToPage = True

    worksheet.sheet_view.showGridLines = (
        settings.show_grid_lines
    )


def export_engineering_original_form(
    definition: EngineeringFormDefinition,
    *,
    survey_record_id,
    file_path,
):
    """
    按 EngineeringFormDefinition.original_form_export_definition
    将一条工程调查记录写入正式 Excel 原表模板。

    模板坐标、评价区、结论区和打印设置
    由 Definition 声明；数据库读取、模板打开、
    顶部归属、写入和保存由本函数统一执行。
    """

    export_definition = (
        definition
        .original_form_export_definition
    )

    if export_definition is None:
        raise ValueError(
            f"{definition.form_code} "
            "尚未配置 OriginalFormExportDefinition。"
        )

    template_path = (
        get_app_root()
        / "templates"
        / "excel"
        / export_definition.template_filename
    )

    if not template_path.exists():
        raise FileNotFoundError(
            "未找到正式原表 Excel 模板。"
            "\n\n"
            f"应存在于：\n{template_path}"
        )

    record = get_engineering_record(
        definition,
        survey_record_id=survey_record_id,
    )

    if record is None:
        raise ValueError(
            "没有找到需要导出的工程调查记录。"
        )

    asset = get_engineering_asset_detail(
        record["engineering_asset_id"]
    )

    if asset is None:
        raise ValueError(
            "没有找到该调查记录对应的工程对象。"
        )

    record_data = (
        record["record_data"]
        or {}
    )

    inspection_results = (
        get_inspection_results(
            survey_record_id
        )
    )

    evaluation_map = {
        result["item_code"]:
        result["grade"]
        for result
        in inspection_results
    }

    workbook = load_workbook(
        template_path
    )

    try:
        if (
            export_definition.sheet_name
            not in workbook.sheetnames
        ):
            raise ValueError(
                "正式原表 Excel 模板中缺少工作表“"
                f"{export_definition.sheet_name}"
                "”。"
            )

        worksheet = workbook[
            export_definition.sheet_name
        ]

        fill_original_form_ownership_header(
            worksheet,
            asset=asset,
            canal_id=(
                _resolve_ownership_canal_id(
                    record
                )
            ),
            business_code=(
                record["business_code"]
            ),
        )

        for cell_binding in (
            export_definition
            .field_bindings
        ):
            value = _resolve_binding(
                record=record,
                record_data=record_data,
                binding=(
                    cell_binding.binding
                ),
            )

            worksheet[
                cell_binding.cell
            ] = (
                ""
                if value is None
                else value
            )

        evaluation_binding = (
            export_definition
            .evaluation_binding
        )

        for row_offset, item in enumerate(
            definition.evaluation_items
        ):
            row_number = (
                evaluation_binding.start_row
                + row_offset
            )

            cell = (
                f"{evaluation_binding.column}"
                f"{row_number}"
            )

            worksheet[cell] = (
                evaluation_map.get(
                    item["item_code"],
                    "",
                )
                or ""
            )

        conclusion = (
            export_definition
            .conclusion_binding
        )

        worksheet[
            conclusion.survey_comment_cell
        ] = (
            record["survey_comment"]
            or ""
        )

        worksheet[
            conclusion.overall_grade_cell
        ] = (
            record["overall_grade"]
            or ""
        )

        worksheet[
            conclusion.survey_date_cell
        ] = (
            record["survey_date"]
            or ""
        )

        _apply_print_settings(
            worksheet,
            export_definition
            .print_settings,
        )

        workbook.save(file_path)

    finally:
        workbook.close()

    return {
        "file_path": file_path,
        "survey_record_id": (
            survey_record_id
        ),
        "business_code": (
            record["business_code"]
        ),
        "asset_name": (
            record["asset_name"]
        ),
    }
