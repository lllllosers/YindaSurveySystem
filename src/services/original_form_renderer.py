"""Database-independent renderer for the desktop and central Web original forms."""

from pathlib import Path

from openpyxl import load_workbook


def render_original_form(
    definition, *, template_root, output, record, record_data,
    evaluation_map, fill_header,
):
    export = definition.original_form_export_definition
    if export is None:
        raise ValueError(f"{definition.form_code} 尚未配置正式原表定义。")
    filename = export.template_filename
    if Path(filename).name != filename or "/" in filename or "\\" in filename or not filename.endswith(".xlsx"):
        raise ValueError("正式原表模板文件名无效。")
    template = Path(template_root) / filename
    if not template.is_file():
        raise FileNotFoundError("未找到正式原表 Excel 模板。")

    workbook = load_workbook(template)
    try:
        if export.sheet_name not in workbook.sheetnames:
            raise ValueError(f"正式原表模板缺少工作表“{export.sheet_name}”。")
        worksheet = workbook[export.sheet_name]
        fill_header(worksheet)
        for cell in export.field_bindings:
            binding = cell.binding
            if binding.source == "record":
                values = [record[key] for key in binding.keys]
            elif binding.source == "record_data":
                values = [record_data.get(key) for key in binding.keys]
            else:
                raise ValueError(f"正式原表不支持的数据源：{binding.source}")
            value = binding.formatter(*values) if binding.formatter else values[0]
            worksheet[cell.cell] = "" if value is None else value

        evaluation = export.evaluation_binding
        for offset, item in enumerate(definition.evaluation_items):
            worksheet[f"{evaluation.column}{evaluation.start_row + offset}"] = (
                evaluation_map.get(item["item_code"]) or ""
            )

        conclusion = export.conclusion_binding
        for attribute, key in (
            ("survey_comment_cell", "survey_comment"),
            ("overall_grade_cell", "overall_grade"),
            ("surveyor_signatures_cell", "surveyor_signatures"),
            ("water_office_manager_signature_cell", "water_office_manager_signature"),
            ("engineering_section_chief_signature_cell", "engineering_section_chief_signature"),
            ("department_head_signature_cell", "department_head_signature"),
            ("survey_date_cell", "survey_date"),
        ):
            worksheet[getattr(conclusion, attribute)] = record.get(key) or ""

        settings = export.print_settings
        worksheet.print_area = settings.print_area
        worksheet.page_setup.orientation = settings.orientation
        if settings.paper_size != "A4":
            raise ValueError(f"通用正式原表导出器暂不支持纸张：{settings.paper_size}")
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = settings.fit_to_width
        worksheet.page_setup.fitToHeight = settings.fit_to_height
        if worksheet.sheet_properties.pageSetUpPr is not None:
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.sheet_view.showGridLines = settings.show_grid_lines
        workbook.save(output)
    finally:
        workbook.close()
