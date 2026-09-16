from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from database import (
    get_app_root,
    get_engineering_asset_detail,
    get_inspection_results,
)

from forms.engineering.form_2_2 import (
    FORM_2_2,
)

from forms.engineering.persistence import (
    get_engineering_record,
)

from services.engineering_summary_export import (
    export_engineering_summary,
)

from services.original_form_export_common import (
    display_value,
    fill_original_form_ownership_header,
)


def export_sluice_gate_summary(
    records,
    file_path,
):
    """
    详细汇总导出的兼容入口。

    实际执行统一交给工程调查通用汇总导出器。
    """

    return export_engineering_summary(
        FORM_2_2,
        records=records,
        file_path=file_path,
    )

def _format_opening_size(record_data):
    """
    将孔数、孔宽、孔高组合为原表要求的：
    孔数/宽×高
    """

    count = display_value(record_data.get("opening_count"))

    width = display_value(record_data.get("opening_width"))

    height = display_value(record_data.get("opening_height"))

    if not count and not width and not height:
        return ""

    if width or height:
        size_text = f"{width}×{height}"
    else:
        size_text = ""

    if count and size_text:
        return f"{count}/{size_text}"

    if count:
        return count

    return size_text


def _get_original_form_template_path():
    """
    获取附表2.2 Excel 模板路径。

    开发环境：
        项目根目录/templates/excel/

    PyInstaller 环境：
        exe 所在目录/templates/excel/
    """

    return get_app_root() / "templates" / "excel" / "form_2_2_V1.xlsx"


def export_sluice_gate_original_form(
    survey_record_id,
    file_path,
):
    """
    将一条附表2.2调查记录填入正式原表模板。
    """

    # =========================
    # 1. 查找模板
    # =========================

    template_path = _get_original_form_template_path()

    if not template_path.exists():
        raise FileNotFoundError(
            "未找到附表2.2 Excel 模板。\n\n" f"应存在于：\n{template_path}"
        )

    # =========================
    # 2. 加载调查记录
    # =========================

    record = get_engineering_record(FORM_2_2, survey_record_id=survey_record_id)

    if record is None:
        raise ValueError("没有找到需要导出的水闸调查记录。")

    asset = get_engineering_asset_detail(record["engineering_asset_id"])

    if asset is None:
        raise ValueError("没有找到该调查记录对应的工程对象。")

    inspection_results = get_inspection_results(survey_record_id)

    evaluation_map = {
        result["item_code"]: result["grade"] for result in inspection_results
    }

    record_data = record["record_data"] or {}

    # =========================
    # 3. 打开模板
    # =========================

    workbook = load_workbook(template_path)
    if "附表2.2" in workbook.sheetnames:
        worksheet = workbook["附表2.2"]
    else:
        worksheet = workbook.active

    if worksheet is None:
        raise ValueError("附表2.2 Excel 模板中没有可用工作表。")

    # =========================
    # 4. 顶部归属和业务编号
    # =========================

    # 原表格式：
    #
    # xxxx 处
    # xxxx 所
    # xxxx 干渠
    # xxxx 支渠
    # 编号：xxxx
    #
    # 因此固定单位文字放在偶数列，
    # 实际名称放在前一个单元格。

    fill_original_form_ownership_header(
        worksheet,
        asset=asset,
        canal_id=record["canal_id"],
        business_code=record["business_code"],
    )

    # =========================
    # 5. 工程基本信息
    # =========================

    # 名称
    worksheet["B5"] = record["asset_name"] or ""

    # 桩号
    worksheet["H5"] = record_data.get("stake") or ""

    # 设计流量
    worksheet["J5"] = record_data.get("design_flow")

    # 建筑物等级
    worksheet["B6"] = record_data.get("structure_grade") or ""

    # 建成年月
    worksheet["D6"] = record_data.get("build_date") or ""

    # 加固改造年月
    worksheet["F6"] = record_data.get("renovation_date") or ""

    # 孔数/宽×高
    worksheet["H6"] = _format_opening_size(record_data)

    # 加大流量
    worksheet["J6"] = record_data.get("increased_flow")

    # 主要构件材料
    worksheet["B7"] = record_data.get("main_component_material") or ""

    # 混凝土强度
    worksheet["D7"] = record_data.get("concrete_strength") or ""

    # 钢筋混凝土强度
    worksheet["F7"] = record_data.get("reinforced_concrete_strength") or ""

    # 保护层厚度
    worksheet["H7"] = record_data.get("cover_thickness")

    # 裂缝限宽
    worksheet["J7"] = record_data.get("crack_width_limit")

    # =========================
    # 6. 14项分项评价
    # =========================
    #
    # 模板第9～22行，
    # 项目类别位于 E 列。

    for row_number, item in enumerate(
        FORM_2_2.evaluation_items,
        start=9,
    ):
        grade = evaluation_map.get(
            item["item_code"],
            "",
        )

        worksheet[f"E{row_number}"] = grade or ""

    # =========================
    # 7. 调查结论
    # =========================

    # 调查意见与建议
    worksheet["C23"] = record["survey_comment"] or ""

    # 工程状况类别
    worksheet["J23"] = record["overall_grade"] or ""

    # 调查时间
    worksheet["J24"] = record["survey_date"] or ""

    # =========================
    # 8. 打印设置
    # =========================

    worksheet.print_area = "A1:J27"

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1

    page_setup_properties = worksheet.sheet_properties.pageSetUpPr

    if page_setup_properties is not None:
        page_setup_properties.fitToPage = True

    worksheet.sheet_view.showGridLines = False

    # =========================
    # 9. 保存
    # =========================

    workbook.save(file_path)

    return {
        "file_path": file_path,
        "survey_record_id": survey_record_id,
        "business_code": record["business_code"],
        "asset_name": record["asset_name"],
    }
